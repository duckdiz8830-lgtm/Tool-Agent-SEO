# -*- coding: utf-8 -*-
"""
SONEALEO ULTRA 3.1 — MÀN HÌNH DỮ LIỆU MỚI (Tab 1 theo thiết kế của Thịnh).

HAI TAB CON, nằm ở FILE RIÊNG — giao-dien.py chỉ gọi gan_vao() 1 dòng,
không phình God Class thêm:

  1. "Làm giàu dữ liệu": quăng file bất kỳ (ảnh / PDF / Excel) vào một
     hàng đợi — tool tự nhận loại, tự moi thông số (biển máy, bảng
     catalogue), kết quả hiện lên BẢNG SỬA TẠI CHỖ (tksheet) với dòng
     nghi vấn tô vàng -> soát -> ghép vào MASTER.
  2. "Xem dữ liệu": mở master Excel, lọc "SKU nào đang thiếu cột gì" —
     trả lời câu 'còn phải đi kiếm data ở đâu'. Kèm nút Mở trong Excel.

ĐÈN TRẠNG THÁI thay cho nút dừng (quyết định của Thịnh: Alt+F4 cho
nhanh): 'Đang đọc — tắt thoải mái' vs 'ĐANG GHI — đừng tắt'. Alt+F4 lúc
ĐỌC an toàn (cache ghi sau từng ảnh); lúc GHI (save Excel) thì ráng chờ
vài giây cho đèn tắt.

tksheet là phụ thuộc MỀM: thiếu thì rơi về bảng chỉ-xem + hướng dẫn cài.
"""
import os
import subprocess
import sys
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from tksheet import Sheet
    CO_TKSHEET = True
except Exception:
    Sheet = None
    CO_TKSHEET = False

DUOI_ANH = ('.jpg', '.jpeg', '.png', '.webp', '.bmp')
MAU_VANG = '#FFF2CC'
MAU_NEN = '#f7f7f5'


def mo_bang_excel(duong_dan):
    """Mở file bằng Excel/LibreOffice của hệ điều hành."""
    try:
        if sys.platform.startswith('win'):
            os.startfile(duong_dan)                      # noqa
        elif sys.platform == 'darwin':
            subprocess.Popen(['open', duong_dan])
        else:
            subprocess.Popen(['xdg-open', duong_dan])
        return True
    except Exception as e:
        messagebox.showerror('Không mở được', f'{e}')
        return False


class DenTrangThai:
    """Đèn báo 'đang đọc / ĐANG GHI' — thay cho nút dừng."""

    def __init__(self, cha):
        self.nhan = tk.Label(cha, text='Sẵn sàng', anchor='w', padx=8,
                             font=('Segoe UI', 9))
        self.nhan.pack(fill='x', pady=(4, 0))
        self.ranh()

    def _dat(self, text, mau_nen, mau_chu):
        def lam():
            self.nhan.config(text=text, bg=mau_nen, fg=mau_chu)
        try:
            self.nhan.after(0, lam)
        except Exception:
            pass

    def dang_doc(self, them=''):
        self._dat(f'● Đang đọc{" — " + them if them else ""} — tắt cửa sổ lúc này KHÔNG sao',
                  '#e8f4ea', '#1a7a2e')

    def dang_ghi(self, them=''):
        self._dat(f'■ ĐANG GHI{" — " + them if them else ""} — ĐỪNG tắt lúc này',
                  '#fdecea', '#b3261e')

    def ranh(self, text='Sẵn sàng'):
        self._dat(text, MAU_NEN, '#666')


# ============================================================ Tab: LÀM GIÀU
class TabLamGiau:
    def __init__(self, cha, app=None):
        self.app = app
        self.f = ttk.Frame(cha)
        self.hang_doi = []          # list (duong_dan, loai)
        self._dung()

    # ---------- dựng UI
    def _dung(self):
        top = ttk.Frame(self.f)
        top.pack(fill='x', padx=10, pady=(10, 4))
        ttk.Label(top, text='Quăng dữ liệu vào — tool tự nhận loại và moi thông số:',
                  font=('Segoe UI', 10, 'bold')).pack(anchor='w')
        nut = ttk.Frame(self.f)
        nut.pack(fill='x', padx=10, pady=4)
        ttk.Button(nut, text='＋ Thêm file...', command=self._them_file).pack(side='left')
        ttk.Button(nut, text='＋ Thêm cả thư mục...', command=self._them_thu_muc)\
            .pack(side='left', padx=6)
        ttk.Button(nut, text='Xoá danh sách', command=self._xoa).pack(side='left')
        ttk.Label(nut, text='   Giới hạn lượt AI:').pack(side='left')
        self.v_gioi_han = tk.StringVar(value='40')
        ttk.Entry(nut, textvariable=self.v_gioi_han, width=5).pack(side='left')
        self.nut_moi = ttk.Button(nut, text='▶ Moi dữ liệu', command=self._moi)
        self.nut_moi.pack(side='left', padx=10)

        self.cay = ttk.Treeview(self.f, columns=('loai', 'tt'), height=5,
                                show='tree headings')
        self.cay.heading('#0', text='File / thư mục')
        self.cay.heading('loai', text='Loại')
        self.cay.heading('tt', text='Trạng thái')
        self.cay.column('#0', width=520)
        self.cay.column('loai', width=110, anchor='center')
        self.cay.column('tt', width=260)
        self.cay.pack(fill='x', padx=10, pady=4)

        self.den = DenTrangThai(self.f)

        # ---------- bảng soát kết quả
        khung_kq = ttk.LabelFrame(self.f, text='  Kết quả — soát và sửa tại chỗ, '
                                               'dòng VÀNG là dòng tool tự thấy nghi  ')
        khung_kq.pack(fill='both', expand=True, padx=10, pady=(6, 10))
        thanh = ttk.Frame(khung_kq)
        thanh.pack(fill='x', pady=4)
        self.v_file_kq = tk.StringVar()
        self.cb_kq = ttk.Combobox(thanh, textvariable=self.v_file_kq, width=40,
                                  state='readonly')
        self.cb_kq.pack(side='left', padx=6)
        ttk.Button(thanh, text='Nạp vào bảng', command=self._nap_bang).pack(side='left')
        ttk.Button(thanh, text='💾 Lưu bảng (ghi đè file kết quả)',
                   command=self._luu_bang).pack(side='left', padx=6)
        ttk.Button(thanh, text='Mở trong Excel', command=self._mo_excel_kq)\
            .pack(side='left')
        ttk.Button(thanh, text='⇒ Ghép vào MASTER...', command=self._ghep)\
            .pack(side='left', padx=6)

        if CO_TKSHEET:
            self.bang = Sheet(khung_kq, header_font=('Segoe UI', 9, 'bold'),
                              font=('Segoe UI', 9, 'normal'), height=260)
            self.bang.enable_bindings('single_select', 'row_select', 'column_width_resize',
                                      'arrowkeys', 'edit_cell', 'copy', 'paste', 'undo',
                                      'rc_select', 'ctrl_click_select', 'right_click_popup_menu')
            self.bang.pack(fill='both', expand=True, padx=4, pady=4)
        else:
            self.bang = None
            ttk.Label(khung_kq, foreground='#b3261e', text=(
                'Chưa cài tksheet nên không sửa tại chỗ được — chạy: '
                'pip install tksheet   (hoặc CAI-DAT-LAN-DAU.bat), rồi mở lại tool.\n'
                'Tạm thời dùng nút "Mở trong Excel" để soát.')).pack(pady=20)
        self._lam_moi_ds_ket_qua()

    # ---------- hàng đợi
    def _loai_file(self, p):
        d = p.lower()
        if os.path.isdir(p):
            return 'thư mục'
        if d.endswith(DUOI_ANH):
            return 'ảnh'
        if d.endswith('.pdf'):
            return 'PDF'
        if d.endswith(('.xlsx', '.xls', '.csv')):
            return 'Excel'
        return 'khác'

    def _them(self, ds):
        for p in ds:
            if not p or any(p == x for x, _l in self.hang_doi):
                continue
            loai = self._loai_file(p)
            self.hang_doi.append((p, loai))
            self.cay.insert('', 'end', text=p, values=(loai, 'chờ'))

    def _them_file(self):
        self._them(list(filedialog.askopenfilenames(
            title='Chọn file (ảnh / PDF / Excel — chọn được nhiều)')))

    def _them_thu_muc(self):
        d = filedialog.askdirectory(title='Chọn thư mục (quét mọi ảnh/PDF bên trong)')
        if d:
            self._them([d])

    def _xoa(self):
        self.hang_doi = []
        for i in self.cay.get_children():
            self.cay.delete(i)

    def _cap_nhat_dong(self, duong_dan, tt):
        def lam():
            for i in self.cay.get_children():
                if self.cay.item(i, 'text') == duong_dan:
                    self.cay.set(i, 'tt', tt)
        try:
            self.f.after(0, lam)
        except Exception:
            pass

    # ---------- moi dữ liệu (chạy nền)
    def _moi(self):
        if not self.hang_doi:
            messagebox.showinfo('Chưa có gì', 'Thêm file hoặc thư mục trước đã.')
            return
        try:
            gh = max(1, int(self.v_gioi_han.get()))
        except ValueError:
            gh = 40
        self.nut_moi.config(state='disabled')
        threading.Thread(target=self._moi_nen, args=(gh,), daemon=True).start()

    def _moi_nen(self, gh):
        import doc_bien_may as B
        import doc_catalogue_pdf as C
        kq_anh, kq_pdf = [], []
        dem = {'anh': 0, 'goi_ai': 0, 'bien_may': 0, 'bo_cache': 0, 'loi': 0}
        try:
            for p, loai in list(self.hang_doi):
                if dem['goi_ai'] >= gh:
                    self._cap_nhat_dong(p, 'bỏ qua — hết giới hạn lượt')
                    continue
                self.den.dang_doc(os.path.basename(p))
                if loai == 'ảnh':
                    self._cap_nhat_dong(p, 'đang đọc...')
                    d = B.doc_mot_anh(p, dem=dem)
                    if d:
                        kq_anh.append(d)
                        self._cap_nhat_dong(p, f'✓ biển máy {d.get("cong_suat_kva") or "?"}kVA'
                                               + (' — CẦN KIỂM' if d.get('can_kiem') else ''))
                    else:
                        self._cap_nhat_dong(p, 'không phải biển máy / lỗi (xem Nhật ký)')
                elif loai == 'thư mục':
                    self._cap_nhat_dong(p, 'đang quét...')
                    # ảnh trong thư mục
                    ket, dem_tm = B.quet_va_doc(p, gioi_han=gh - dem['goi_ai'])
                    kq_anh += ket
                    for k in dem:
                        dem[k] += dem_tm.get(k, 0)
                    # PDF trong thư mục
                    for goc_tm, _d, fs in os.walk(p):
                        for f in sorted(fs):
                            if f.lower().endswith('.pdf') and dem['goi_ai'] < gh:
                                dong, n = C.doc_mot_pdf(os.path.join(goc_tm, f),
                                                        gh - dem['goi_ai'])
                                kq_pdf += dong
                                dem['goi_ai'] += n
                    self._cap_nhat_dong(p, f'✓ {len(ket)} biển máy + PDF bên trong')
                elif loai == 'PDF':
                    self._cap_nhat_dong(p, 'đang đọc...')
                    dong, n = C.doc_mot_pdf(p, gh - dem['goi_ai'])
                    kq_pdf += dong
                    dem['goi_ai'] += n
                    self._cap_nhat_dong(p, f'✓ {len(dong)} mức công suất')
                elif loai == 'Excel':
                    self._cap_nhat_dong(p, 'Excel: dùng nút "Ghép vào MASTER" nếu là '
                                           'file kết quả đã soát')
                else:
                    self._cap_nhat_dong(p, 'loại này chưa hỗ trợ tự động')
            # ghi file kết quả
            if kq_anh:
                self.den.dang_ghi('bien-may-ket-qua.xlsx')
                B.xuat_excel(kq_anh)
            if kq_pdf:
                self.den.dang_ghi('catalogue-ket-qua.xlsx')
                C.xuat_excel(kq_pdf)
            self.den.ranh(f'Xong — {dem["goi_ai"]} lượt AI, {len(kq_anh)} biển máy, '
                          f'{len(kq_pdf)} dòng catalogue, cache đỡ {dem["bo_cache"]}')
            print(f'== [Làm giàu 3.1] {len(kq_anh)} biển máy + {len(kq_pdf)} dòng catalogue, '
                  f'{dem["goi_ai"]} lượt AI, lỗi {dem["loi"]} ==')
            self.f.after(0, self._lam_moi_ds_ket_qua)
            if kq_anh or kq_pdf:
                self.f.after(0, lambda: messagebox.showinfo(
                    'Moi xong', 'Đã có file kết quả — nạp vào bảng bên dưới để soát.\n'
                    'Dòng VÀNG là dòng tool tự thấy nghi: đúng thì xoá dấu X ở cột '
                    '"Cần kiểm", không chắc thì để nguyên (sẽ bị bỏ khi ghép).'))
        except Exception as e:
            self.den.ranh(f'Lỗi: {e}')
            print(f'[Làm giàu 3.1] LỖI: {e}')
        finally:
            self.f.after(0, lambda: self.nut_moi.config(state='normal'))

    # ---------- bảng soát
    def _lam_moi_ds_ket_qua(self):
        co = [f for f in ('bien-may-ket-qua.xlsx', 'catalogue-ket-qua.xlsx')
              if os.path.exists(f)]
        self.cb_kq['values'] = co
        if co and not self.v_file_kq.get():
            self.v_file_kq.set(co[0])

    def _doc_xlsx(self, duong_dan):
        import openpyxl
        wb = openpyxl.load_workbook(duong_dan, read_only=True)
        ws = wb.active
        rows = [[('' if c is None else c) for c in r]
                for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows

    def _nap_bang(self):
        f = self.v_file_kq.get()
        if not f or not os.path.exists(f):
            messagebox.showinfo('Chưa có file', 'Chưa có file kết quả nào — bấm "Moi dữ liệu" trước.')
            return
        rows = self._doc_xlsx(f)
        if not rows:
            return
        if not self.bang:
            mo_bang_excel(f)
            return
        hdr, data = rows[0], rows[1:]
        self.bang.set_sheet_data([[str(c) for c in r] for r in data],
                                 reset_col_positions=True, reset_row_positions=True)
        self.bang.headers([str(h) for h in hdr])
        # tô vàng dòng Cần kiểm = X
        try:
            i_ck = hdr.index('Cần kiểm')
            self.bang.dehighlight_all()
            for i, r in enumerate(data):
                if str(r[i_ck]).strip().upper() == 'X':
                    self.bang.highlight_rows(rows=[i], bg=MAU_VANG, fg='black')
        except ValueError:
            pass
        self.bang.set_all_cell_sizes_to_text()

    def _luu_bang(self):
        f = self.v_file_kq.get()
        if not f or not self.bang:
            return
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        self.den.dang_ghi(f)
        try:
            hdr = [str(h) for h in self.bang.headers()]
            data = self.bang.get_sheet_data()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(hdr)
            for c in ws[1]:
                c.font = Font(bold=True)
            vang = PatternFill('solid', fgColor='FFF2CC')
            i_ck = hdr.index('Cần kiểm') if 'Cần kiểm' in hdr else -1
            for r in data:
                ws.append(list(r))
                if i_ck >= 0 and str(r[i_ck]).strip().upper() == 'X':
                    for c in ws[ws.max_row]:
                        c.fill = vang
            wb.save(f)
            self.den.ranh(f'Đã lưu {f}')
        except Exception as e:
            self.den.ranh(f'Lỗi lưu: {e}')
            messagebox.showerror('Không lưu được', str(e))

    def _mo_excel_kq(self):
        f = self.v_file_kq.get()
        if f and os.path.exists(f):
            mo_bang_excel(os.path.abspath(f))

    def _ghep(self):
        master = filedialog.askopenfilename(title='Chọn file MASTER cần ghép vào',
                                            filetypes=[('Excel', '*.xlsx')])
        if not master:
            return
        def chay():
            import ghep_vao_master as G
            self.den.dang_ghi('ghép vào bản sao MASTER')
            try:
                ra = G.ghep(master, [('bien-may-ket-qua.xlsx', 0),
                                     ('catalogue-ket-qua.xlsx', 1)])
                self.den.ranh('Ghép xong' if ra else 'Không có dòng nào dùng được')
                if ra:
                    self.f.after(0, lambda: messagebox.showinfo(
                        'Ghép xong', f'Đã tạo:\n{ra}\n\nMASTER gốc KHÔNG bị sửa. '
                        f'Xem sheet "DoiChieu" nếu có số lệch nhau.\n'
                        f'Nhập file này ở tab "Nhập dữ liệu" rồi chạy thử 5-10 bài.'))
                else:
                    self.f.after(0, lambda: messagebox.showinfo(
                        'Chưa ghép được', 'Mọi dòng còn dấu X "Cần kiểm" (hoặc thiếu kVA) '
                        'đều bị bỏ qua — soát bảng rồi thử lại.'))
            except Exception as e:
                self.den.ranh(f'Lỗi ghép: {e}')
                print(f'[Ghép 3.1] LỖI: {e}')
        threading.Thread(target=chay, daemon=True).start()


# ============================================================ Tab: XEM DỮ LIỆU
class TabXemDuLieu:
    def __init__(self, cha, app=None):
        self.f = ttk.Frame(cha)
        self.duong_dan = None
        self.hdr, self.data = [], []
        self._dung()

    def _dung(self):
        thanh = ttk.Frame(self.f)
        thanh.pack(fill='x', padx=10, pady=(10, 4))
        ttk.Button(thanh, text='Chọn file Excel...', command=self._chon).pack(side='left')
        ttk.Label(thanh, text=' Sheet:').pack(side='left')
        self.v_sheet = tk.StringVar()
        self.cb_sheet = ttk.Combobox(thanh, textvariable=self.v_sheet, width=26,
                                     state='readonly')
        self.cb_sheet.pack(side='left', padx=4)
        self.cb_sheet.bind('<<ComboboxSelected>>', lambda e: self._nap())
        ttk.Label(thanh, text=' Soi cột còn trống:').pack(side='left')
        self.v_cot = tk.StringVar(value='(tất cả dòng)')
        self.cb_cot = ttk.Combobox(thanh, textvariable=self.v_cot, width=34,
                                   state='readonly')
        self.cb_cot.pack(side='left', padx=4)
        self.cb_cot.bind('<<ComboboxSelected>>', lambda e: self._loc())
        ttk.Button(thanh, text='Mở trong Excel', command=self._mo_ngoai).pack(side='left', padx=6)
        ttk.Button(thanh, text='⟳ Nạp lại', command=self._nap).pack(side='left')

        self.nhan_tk = ttk.Label(self.f, text='Chưa mở file nào.', padding=(10, 2))
        self.nhan_tk.pack(fill='x')

        if CO_TKSHEET:
            self.bang = Sheet(self.f, header_font=('Segoe UI', 9, 'bold'),
                              font=('Segoe UI', 9, 'normal'))
            self.bang.enable_bindings('single_select', 'row_select', 'column_width_resize',
                                      'arrowkeys', 'copy', 'right_click_popup_menu')
            self.bang.pack(fill='both', expand=True, padx=10, pady=(4, 10))
        else:
            self.bang = None
            ttk.Label(self.f, foreground='#b3261e',
                      text='Chưa cài tksheet — pip install tksheet rồi mở lại tool.\n'
                           'Tạm dùng nút "Mở trong Excel".').pack(pady=30)

    def _chon(self):
        f = filedialog.askopenfilename(title='Chọn file Excel',
                                       filetypes=[('Excel', '*.xlsx')])
        if not f:
            return
        self.duong_dan = f
        import openpyxl
        wb = openpyxl.load_workbook(f, read_only=True)
        self.cb_sheet['values'] = wb.sheetnames
        # ưu tiên sheet có dữ liệu sản phẩm
        chon = next((s for s in wb.sheetnames if not s.startswith('00')), wb.sheetnames[0])
        wb.close()
        self.v_sheet.set(chon)
        self._nap()

    def _nap(self):
        if not self.duong_dan:
            return
        import openpyxl
        wb = openpyxl.load_workbook(self.duong_dan, read_only=True)
        ws = wb[self.v_sheet.get()] if self.v_sheet.get() in wb.sheetnames else wb.active
        rows = [[('' if c is None else c) for c in r]
                for r in ws.iter_rows(values_only=True)]
        wb.close()
        if not rows:
            return
        self.hdr = [str(h) for h in rows[0]]
        self.data = rows[1:]
        self.cb_cot['values'] = ['(tất cả dòng)'] + self.hdr
        self.v_cot.set('(tất cả dòng)')
        self._loc()

    def _loc(self):
        if not self.hdr:
            return
        cot = self.v_cot.get()
        if cot == '(tất cả dòng)':
            loc = self.data
            self.nhan_tk.config(text=f'{len(loc)} dòng — chọn 1 cột để soi dòng còn trống.')
        else:
            i = self.hdr.index(cot)
            loc = [r for r in self.data
                   if i >= len(r) or str(r[i]).strip() in ('', '-', 'None')]
            self.nhan_tk.config(
                text=f'{len(loc)}/{len(self.data)} dòng ĐANG TRỐNG cột "{cot}" — '
                     f'đây là chỗ cần đi kiếm data tiếp.')
        if self.bang:
            gioi_han = loc[:800]      # đủ soi, không nghẽn UI
            self.bang.set_sheet_data([[str(c) for c in r] for r in gioi_han],
                                     reset_col_positions=True, reset_row_positions=True)
            self.bang.headers(self.hdr)

    def _mo_ngoai(self):
        if self.duong_dan:
            mo_bang_excel(self.duong_dan)


# ============================================================ gắn vào tool
def gan_vao(nb_con, app=None):
    """giao-dien.py gọi đúng 1 dòng: man_hinh_du_lieu.gan_vao(nb, self)."""
    t1 = TabLamGiau(nb_con, app)
    nb_con.add(t1.f, text='  Làm giàu dữ liệu (3.1)  ')
    t2 = TabXemDuLieu(nb_con, app)
    nb_con.add(t2.f, text='  Xem dữ liệu (3.1)  ')
    return t1, t2
