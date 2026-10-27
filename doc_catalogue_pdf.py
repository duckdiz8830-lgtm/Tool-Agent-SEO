# -*- coding: utf-8 -*-
"""
SONEALEO ULTRA 3.0 — ĐỌC CATALOGUE PDF: tự tìm trang có BẢNG THÔNG SỐ
theo kVA trong catalogue, đọc bảng thành Excel để người soát.

VÌ SAO KHÁC LẦN THẤT BẠI CŨ (12 PDF -> 15/245 SKU)
Lần trước trích TEXT: bảng vỡ cấu trúc, số của 250kVA dính vào số của
400kVA. Lần này render trang thành ẢNH rồi cho AI vision đọc — bảng
còn nguyên hàng nguyên cột, một trang bảng của một hãng lấp được cả
dải công suất (15-20 SKU) một lượt.

VAI NGƯỜI - VAI MÁY (người dùng KHÔNG cần đọc catalogue):
  Máy: bung PDF thành trang, tự chấm trang nào có bảng thông số (lọc
  nhanh bằng text layer trước — 0 quota; chỉ trang khả nghi mới hỏi
  vision), đọc bảng, xuất Excel ghi rõ 'file X trang Y'.
  Người: mở Excel, liếc số nào vô lý thì mở đúng trang đã ghi nguồn.

CÁCH DÙNG
    python doc_catalogue_pdf.py "thu-muc-pdf-hoac-1-file.pdf" [gioi_han_goi_ai]
"""
import os
import re
import sys

import vision_chung as V

# từ khoá gợi ý trang bảng thông số MBA (đọc từ text layer — miễn phí)
_TU_KHOA = ['kva', 'po', 'pk', 'uk', 'tổn hao', 'ton hao', 'no-load',
            'no load', 'load loss', 'short circuit', 'impedance', 'dyn',
            'khối lượng', 'khoi luong', 'weight', 'kích thước', 'dimension']

PROMPT_PHAN_LOAI = (
    'Trang catalogue này có chứa BẢNG THÔNG SỐ KỸ THUẬT máy biến áp theo '
    'từng mức công suất kVA không (mỗi hàng/cột một mức kVA, có các thông '
    'số như Po, Pk, Uk, kích thước, khối lượng)? Trang chỉ có hình minh '
    'hoạ, giới thiệu công ty, mục lục -> false. '
    'Trả JSON: {"co_bang_thong_so": true/false}'
)

PROMPT_DOC_BANG = (
    'Trang này chứa bảng thông số máy biến áp theo mức công suất. CHỈ CHÉP '
    'các giá trị in trong bảng, ô trống hoặc không đọc rõ -> null. Không '
    'suy đoán, không tính toán, không nội suy.\n'
    'Trả JSON: {"hang": "tên hãng nếu in trên trang"|null, '
    '"loai_may": "dau|kho|null", '
    '"dien_ap": "cấp điện áp chung của bảng nếu có"|null, '
    '"cac_dong": [{"cong_suat_kva": số, "po_w": số|null, "pk_w": số|null, '
    '"uk_pct": số|null, "io_pct": số|null, "dien_ap": "..."|null, '
    '"to_dau_day": "..."|null, "tong_khoi_luong_kg": số|null, '
    '"khoi_luong_dau_kg": số|null, "kich_thuoc_mm": "DxRxC"|null}], '
    '"do_ro": "cao|vua|thap"}\n'
    'LƯU Ý ĐƠN VỊ: nếu bảng ghi tổn hao bằng kW thì đổi ra W (nhân 1000) '
    'và chỉ khi đơn vị in rõ trên bảng; không rõ đơn vị -> null.'
)

COT_XUAT = [
    ('nguon', 'Nguồn (file / trang)'), ('hang', 'Hãng'), ('loai_may', 'Dầu/Khô'),
    ('cong_suat_kva', 'Công suất (kVA)'), ('dien_ap', 'Điện áp'),
    ('to_dau_day', 'Tổ đấu dây'), ('po_w', 'Tổn hao không tải Po (W)'),
    ('pk_w', 'Tổn hao ngắn mạch Pk (W)'), ('uk_pct', 'Điện áp ngắn mạch Uk (%)'),
    ('io_pct', 'Dòng không tải Io (%)'), ('khoi_luong_dau_kg', 'Khối lượng dầu (kg)'),
    ('tong_khoi_luong_kg', 'Tổng khối lượng (kg)'), ('kich_thuoc_mm', 'Kích thước (mm)'),
    ('do_ro', 'Độ rõ'), ('can_kiem', 'Cần kiểm'),
]


def _diem_trang(text):
    """Chấm điểm text layer của trang: càng nhiều từ khoá + con số càng
    giống trang bảng thông số. 0 quota."""
    t = (text or '').lower()
    diem = sum(1 for k in _TU_KHOA if k in t)
    diem += min(len(re.findall(r'\d{2,}', t)) // 20, 5)
    return diem


def trang_ung_vien(pdf_path, toi_da=6):
    """Trả list (số trang, điểm) đáng gửi vision. PDF scan (không text
    layer) -> trả mọi trang, vision phân loại thay."""
    import fitz
    doc = fitz.open(pdf_path)
    cham = []
    co_text = False
    for i, trang in enumerate(doc):
        text = trang.get_text() or ''
        if len(text.strip()) > 50:
            co_text = True
        cham.append((i, _diem_trang(text)))
    doc.close()
    if co_text:
        cham = [c for c in cham if c[1] >= 3]
        cham.sort(key=lambda x: -x[1])
        return cham[:toi_da], True
    return [(i, 0) for i, _ in cham], False


def _trang_thanh_anh(pdf_path, so_trang, dpi=150):
    import fitz
    doc = fitz.open(pdf_path)
    pix = doc[so_trang].get_pixmap(dpi=dpi)
    anh = pix.tobytes('jpeg')
    doc.close()
    return anh, 'image/jpeg'


def doc_mot_pdf(pdf_path, gioi_han_con_lai, goi_vision=None):
    goi = goi_vision or V.goi_vision
    ch = V.doc_cau_hinh()
    ten = os.path.basename(pdf_path)
    ung_vien, co_text = trang_ung_vien(pdf_path)
    if not co_text:
        print(f'  {ten}: PDF scan (không có text layer) — vision sẽ tự lọc trang, tốn lượt hơn.')
    _k_pl = ch['model'] + PROMPT_PHAN_LOAI      # Ver 3.1 — cache theo cả model+prompt
    _k_doc = ch['model'] + PROMPT_DOC_BANG
    ket_qua, goi_ai = [], 0
    for so_trang, diem in ung_vien:
        nhan = f'pdf-t{so_trang}'
        # cache theo (file, trang) — nhét số trang vào nhãn cache
        da = V.doc_cache(pdf_path, nhan + '-doc', them=_k_doc)
        if da is None:
            if goi_ai >= gioi_han_con_lai:
                break
            anh, mime = _trang_thanh_anh(pdf_path, so_trang)
            pl = V.doc_cache(pdf_path, nhan + '-pl', them=_k_pl)
            if pl is None:
                kq = goi(anh, mime, PROMPT_PHAN_LOAI, cau_hinh=ch, max_tokens=1200)
                goi_ai += 1
                if kq['ok']:
                    pl = kq['json']
                    V.ghi_cache(pdf_path, nhan + '-pl', pl, them=_k_pl)
                else:
                    # Ver 3.1 — AI chết thì FALLBACK RULE: trang có text layer
                    # điểm từ khoá cao (>=5) coi như có bảng, đi tiếp bước đọc
                    # thay vì vứt trang (yêu cầu GPT: 'fallback bằng rule').
                    if co_text and diem >= 5:
                        pl = {'co_bang_thong_so': True}
                        print(f'  [phân loại AI lỗi -> dùng rule] {ten} trang {so_trang + 1} '
                              f'(điểm từ khoá {diem}): {kq["loi"]}')
                    else:
                        print(f'  [LỖI phân loại] {ten} trang {so_trang + 1}: {kq["loi"]}')
                        continue
            if not pl.get('co_bang_thong_so'):
                continue
            if goi_ai >= gioi_han_con_lai:
                break
            kq = goi(anh, mime, PROMPT_DOC_BANG, cau_hinh=ch, max_tokens=6000)
            goi_ai += 1
            if not kq['ok']:
                print(f'  [LỖI đọc bảng] {ten} trang {so_trang + 1}: {kq["loi"]}')
                continue
            da = kq['json']
            V.ghi_cache(pdf_path, nhan + '-doc', da, them=_k_doc)
        hang, loai = da.get('hang'), da.get('loai_may')
        for dong in da.get('cac_dong') or []:
            d = dict(dong)
            d.setdefault('hang', hang)
            d.setdefault('loai_may', loai)
            d.setdefault('dien_ap', dong.get('dien_ap') or da.get('dien_ap'))
            d['nguon'] = f'{ten} / trang {so_trang + 1}'
            d['do_ro'] = da.get('do_ro')
            d['can_kiem'] = 'X' if da.get('do_ro') == 'thap' else ''
            ket_qua.append(d)
        print(f'  [OK] {ten} trang {so_trang + 1}: {len(da.get("cac_dong") or [])} mức công suất')
    return ket_qua, goi_ai


def xuat_excel(ket_qua, duong_dan_ra='catalogue-ket-qua.xlsx'):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Catalogue'
    ws.append([nhan for _k, nhan in COT_XUAT])
    for c in ws[1]:
        c.font = Font(bold=True)
    vang = PatternFill('solid', fgColor='FFF2CC')
    for d in ket_qua:
        ws.append([d.get(k, '') for k, _n in COT_XUAT])
        if d.get('can_kiem'):
            for c in ws[ws.max_row]:
                c.fill = vang
    wb.save(duong_dan_ra)
    return duong_dan_ra


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Cách dùng: python doc_catalogue_pdf.py <thư mục PDF hoặc 1 file .pdf> [giới hạn lượt AI]')
        sys.exit(1)
    nguon = sys.argv[1]
    gioi_han = int(sys.argv[2]) if len(sys.argv) > 2 else 60
    ds_pdf = ([nguon] if nguon.lower().endswith('.pdf')
              else [os.path.join(nguon, f) for f in sorted(os.listdir(nguon))
                    if f.lower().endswith('.pdf')])
    tat_ca, da_goi = [], 0
    for p in ds_pdf:
        kq, n = doc_mot_pdf(p, gioi_han - da_goi)
        tat_ca += kq
        da_goi += n
        if da_goi >= gioi_han:
            print('Đã chạm giới hạn lượt AI — chạy lại sau, cache sẽ đọc tiếp phần còn thiếu.')
            break
    if tat_ca:
        ra = xuat_excel(tat_ca)
        print(f'\nXong: {len(tat_ca)} dòng công suất từ {len(ds_pdf)} file, {da_goi} lượt AI.')
        print(f'Kết quả: {ra} — MỞ FILE, soát dòng tô vàng trước khi ghép.')
    else:
        print(f'\nKhông đọc được bảng nào ({da_goi} lượt AI đã dùng).')
