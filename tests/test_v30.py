# -*- coding: utf-8 -*-
"""Ver 3.0 — test: parser đường dẫn, pipeline vision (mock), PDF, ghép master.

Chạy: python tests/test_v30.py  (từ thư mục v30). KHÔNG gọi mạng thật.
"""
import json
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import anh_san_pham as A
import vision_chung as V
import doc_bien_may as B
import doc_catalogue_pdf as C
import ghep_vao_master as G

PASS = FAIL = 0


def kiem(dk, ten):
    global PASS, FAIL
    PASS, FAIL = (PASS + 1, FAIL) if dk else (PASS, FAIL + 1)
    print(f'  {"PASS" if dk else "FAIL"}  {ten}')


# ================================================= 1. parser đường dẫn
print('== 1. Đường dẫn thư mục là metadata ==')
with tempfile.TemporaryDirectory() as tmp:
    for p in ['MBA/320', 'MBA/1250 - 250', 'MBA/HBT', 'MBA/2000/bienmay',
              'CAP/320']:
        os.makedirs(os.path.join(tmp, p), exist_ok=True)
    duong = {
        'vo_danh_320': os.path.join(tmp, 'MBA/320/z4395168565850_0.jpg'),
        'mo_ho': os.path.join(tmp, 'MBA/1250 - 250/z111.jpg'),
        'hbt': os.path.join(tmp, 'MBA/HBT/z222.jpg'),
        'bien_may': os.path.join(tmp, 'MBA/2000/bienmay/z333.jpg'),
        'cap_320': os.path.join(tmp, 'CAP/320/z444.jpg'),
        'ten_thang': os.path.join(tmp, 'MBA/320/mba dau thibidi 3p 560kva.jpg'),
    }
    for f in duong.values():
        open(f, 'wb').write(b'x')

    t = A.phan_tich_anh(duong['vo_danh_320'], tmp)
    kiem(t['cong_suat'] == 320 and 'thư mục' in t['nguon_path'],
         "file vô danh trong MBA/320/ -> 320kVA (nguồn: thư mục)")
    kiem(t['do_tin_cay'] == 'vua', "kVA từ thư mục -> tin cậy 'vừa' (trần), không 'cao'")

    t = A.phan_tich_anh(duong['mo_ho'], tmp)
    kiem(t['cong_suat'] is None and t['canh_bao'] and '1250' in t['canh_bao'],
         "thư mục '1250 - 250' (2 số) -> KHÔNG đoán, có cảnh báo")

    t = A.phan_tich_anh(duong['hbt'], tmp)
    kiem(t['thuong_hieu'] == 'HBT', 'thư mục HBT -> hãng HBT')

    t = A.phan_tich_anh(duong['bien_may'], tmp)
    kiem(t['la_bien_may'] is True, "thư mục bienmay/ -> đánh dấu ảnh biển máy")
    kiem(t['cong_suat'] == 2000, 'ảnh biển vẫn mang kVA 2000 từ thư mục cha')

    t = A.phan_tich_anh(duong['cap_320'], tmp)
    kiem(t['cong_suat'] is None, "folder '320' KHÔNG có ngữ cảnh MBA -> không nhận là kVA")

    t = A.phan_tich_anh(duong['ten_thang'], tmp)
    kiem(t['cong_suat'] == 560 and t['thuong_hieu'] == 'THIBIDI',
         'tên file ĐẦY ĐỦ thắng thư mục (560kVA THIBIDI, không phải 320)')
    kiem(t['do_tin_cay'] == 'cao', 'tên file đầy đủ giữ tin cậy cao như cũ')

    # quet_thu_muc: ảnh biển máy không được ghép SKU
    rows = [{'ma_goc': 'X320', 'ten_goc': 'Máy biến áp dầu 3Pha 320KVA MBT',
             'cong_suat': '320', 'thuong_hieu': 'MBT', 'record_id': 1}]
    kq = A.quet_thu_muc(tmp, rows)
    d_bm = next(d for d in kq['danh_sach'] if 'z333' in d['ten_file'])
    kiem(d_bm['chon'] is False and 'biển máy' in d_bm['ly_do'],
         'quet_thu_muc: ảnh biển máy không ghép, không tick, lý do rõ')
    d_vd = next(d for d in kq['danh_sach'] if 'z4395168565850' in d['ten_file'])
    kiem(d_vd['ma_goc'] == 'X320' and d_vd['chon'] is False,
         'ảnh vô danh ĐƯỢC đề xuất SKU 320 nhưng KHÔNG tự tick (tin=vừa)')

# ================================================= 2. vision mock — biển máy
print('== 2. Đọc biển máy (vision mock, không mạng) ==')


def mock_vision_factory(tra_loi_theo_prompt):
    def mock(anh_bytes, mime, prompt, cau_hinh=None, max_tokens=0, **kw):
        for khoa, tra in tra_loi_theo_prompt.items():
            if khoa in prompt:
                return {'ok': True, 'json': tra}
        return {'ok': False, 'loi': 'mock: prompt lạ'}
    return mock


with tempfile.TemporaryDirectory() as tmp:
    os.makedirs(os.path.join(tmp, 'MBA/320'))
    # tạo ảnh JPEG thật (PIL cần mở được)
    from PIL import Image
    Image.new('RGB', (800, 600), (200, 200, 200)).save(
        os.path.join(tmp, 'MBA/320', 'bien.jpg'))
    Image.new('RGB', (801, 600), (90, 120, 150)).save(
        os.path.join(tmp, 'MBA/320', 'may_tren_xe.jpg'))
    cwd = os.getcwd()
    os.chdir(tmp)      # cache ghi tại chỗ, không bẩn thư mục dự án
    try:
        mock = mock_vision_factory({
            'BIỂN THÔNG SỐ': None,      # điền bên dưới theo từng file — đơn giản: phân loại theo tên không được với mock chung, nên trả true cho cả 2 rồi kiểm đối chiếu
        })
        # mock tinh hơn: phân loại true cho 'bien.jpg' — nhận diện qua bytes không được,
        # nên dùng 2 lượt: lượt phân loại luôn true, lượt đọc trả kVA lệch cho 1 case
        goi_dem = {'n': 0}

        def mock2(anh_bytes, mime, prompt, cau_hinh=None, max_tokens=0, **kw):
            if 'BIỂN THÔNG SỐ' in prompt or 'nameplate' in prompt:
                return {'ok': True, 'json': {'la_bien_may': True}}
            goi_dem['n'] += 1
            if goi_dem['n'] == 1:
                return {'ok': True, 'json': {
                    'thuong_hieu': 'THIBIDI', 'loai_may': 'dau', 'so_pha': 3,
                    'cong_suat_kva': 320, 'dien_ap': '22/0.4 kV',
                    'to_dau_day': 'Dyn11', 'po_w': 510, 'pk_w': 3400,
                    'uk_pct': 4.5, 'io_pct': None, 'khoi_luong_dau_kg': 260,
                    'tong_khoi_luong_kg': 1180, 'kich_thuoc_mm': None,
                    'tieu_chuan': 'TCVN 6306-1', 'nam_sx': 2024, 'do_ro': 'cao'}}
            return {'ok': True, 'json': {
                'thuong_hieu': 'MBT', 'loai_may': 'dau', 'so_pha': 3,
                'cong_suat_kva': 560, 'po_w': 700, 'pk_w': None, 'uk_pct': None,
                'io_pct': None, 'khoi_luong_dau_kg': None,
                'tong_khoi_luong_kg': None, 'kich_thuoc_mm': None,
                'dien_ap': None, 'to_dau_day': None,
                'tieu_chuan': None, 'nam_sx': None, 'do_ro': 'cao'}}

        kq, dem = B.quet_va_doc('MBA', gioi_han=99, goi_vision=mock2)
        kiem(len(kq) == 2 and dem['bien_may'] == 2, 'đọc được 2 biển máy (mock)')
        d_khop = next(d for d in kq if d.get('cong_suat_kva') == 320)
        kiem(d_khop['can_kiem'] == '' and d_khop['po_w'] == 510,
             'biển 320kVA khớp thư mục 320 -> KHÔNG cần kiểm, Po=510W')
        d_lech = next(d for d in kq if d.get('cong_suat_kva') == 560)
        kiem(d_lech['can_kiem'] == 'X' and 'thư mục 320' in d_lech['ly_do_kiem'],
             'biển 560kVA nằm trong thư mục 320 -> ĐÁNH DẤU cần kiểm')
        ra = B.xuat_excel(kq, 'bien-may-ket-qua.xlsx')
        kiem(os.path.exists(ra), 'xuất Excel kết quả biển máy')

        # cache: chạy lại không gọi AI lượt nào
        truoc = goi_dem['n']
        kq2, dem2 = B.quet_va_doc('MBA', gioi_han=99, goi_vision=mock2)
        kiem(goi_dem['n'] == truoc and dem2['goi_ai'] == 0,
             'chạy lại -> cache đỡ toàn bộ, 0 lượt AI')

        # ============================================ 3. PDF (mock vision)
        print('== 3. Catalogue PDF ==')
        import fitz
        doc = fitz.open()
        t1 = doc.new_page()
        t1.insert_text((50, 80), 'GIOI THIEU CONG TY ABC ...')
        t2 = doc.new_page()
        t2.insert_text((50, 60), 'Thong so ky thuat  kVA  Po  Pk  Uk  ton hao  '
                                 '100 250 320 400 560 630 750 800 1000 1250 '
                                 'khoi luong kich thuoc Dyn11 22/0.4')
        doc.save('catalog_test.pdf')
        doc.close()
        uv, co_text = C.trang_ung_vien('catalog_test.pdf')
        kiem(co_text and len(uv) == 1 and uv[0][0] == 1,
             'lọc text layer 0-quota: chỉ trang 2 (có bảng) vào vòng vision')

        def mock_pdf(anh_bytes, mime, prompt, cau_hinh=None, max_tokens=0, **kw):
            if 'co_bang_thong_so' in prompt:
                return {'ok': True, 'json': {'co_bang_thong_so': True}}
            return {'ok': True, 'json': {
                'hang': 'THIBIDI', 'loai_may': 'dau', 'dien_ap': '22/0.4 kV',
                'cac_dong': [
                    {'cong_suat_kva': 250, 'po_w': 440, 'pk_w': 2750, 'uk_pct': 4.0},
                    {'cong_suat_kva': 320, 'po_w': 510, 'pk_w': 3400, 'uk_pct': 4.5},
                ], 'do_ro': 'cao'}}

        kq_pdf, n = C.doc_mot_pdf('catalog_test.pdf', 99, goi_vision=mock_pdf)
        kiem(len(kq_pdf) == 2 and kq_pdf[0]['hang'] == 'THIBIDI',
             '1 trang bảng -> 2 dòng công suất, hãng lan xuống từng dòng')
        kiem('trang 2' in kq_pdf[0]['nguon'], 'từng dòng ghi rõ nguồn file/trang')
        C.xuat_excel(kq_pdf, 'catalogue-ket-qua.xlsx')

        # ============================================ 4. ghép vào master
        print('== 4. Ghép vào master — không ghi đè, xung đột báo người ==')
        import openpyxl
        wbm = openpyxl.Workbook()
        wsm = wbm.active
        wsm.title = '03_May_Bien_Ap_Do_Luong'
        wsm.append(['Mã hàng', 'Tên hàng', 'Nhãn Hiệu / Xuất Xứ', 'Công Suất (kVA)',
                    'Tổn hao không tải Po (W, catalogue)'])
        wsm.append(['TBD3P320', 'Máy biến áp dầu 3Pha 320KVA 22/0.4KV - TBD',
                    'THIBIDI', '320 kVA', ''])          # trống -> được điền
        wsm.append(['TBD3P250', 'Máy biến áp dầu 3Pha 250KVA 22/0.4KV - TBD',
                    'THIBIDI', '250 kVA', 999])         # có sẵn 999 ≠ 440 -> xung đột
        wsm.append(['KPK320', 'Máy biến áp Khô KP 320KVA', 'KP', '320 kVA', ''])
        wbm.save('master_test.xlsx')

        ra = G.ghep('master_test.xlsx',
                    [('bien-may-ket-qua.xlsx', 0), ('catalogue-ket-qua.xlsx', 1)])
        wb2 = openpyxl.load_workbook(ra)
        ws2 = wb2['03_May_Bien_Ap_Do_Luong']
        hdr = {c.value: i + 1 for i, c in enumerate(ws2[1]) if c.value}
        po_col = hdr['Tổn hao không tải Po (W) [đọc từ ảnh]']
        kiem(ws2.cell(row=2, column=po_col).value == 510,
             'TBD 320 dầu: điền Po=510 (biển máy ưu tiên, trùng catalogue)')
        kiem(ws2.cell(row=4, column=po_col).value in (None, ''),
             'KP KHÔ 320: KHÔNG điền — hãng và dầu/khô không khớp nguồn')
        kiem(ws2.cell(row=3, column=hdr['Tổn hao không tải Po (W, catalogue)']).value == 999,
             'giá trị sẵn có (999) KHÔNG bị ghi đè')
        ds_dc = [r for r in wb2['DoiChieu'].iter_rows(values_only=True)]
        kiem(any('TBD3P250' in str(r[0]) or '250KVA' in str(r[0]) for r in ds_dc[1:]),
             'lệch 440 vs 999 -> nằm trong sheet DoiChieu chờ người quyết')
        # dòng "Cần kiểm" còn dấu X không được dùng
        nguon_lech = [d for d in G._doc_ket_qua('bien-may-ket-qua.xlsx', 0)]
        kiem(all(d.get('cong_suat_kva') != 560 for d in nguon_lech),
             'dòng biển máy còn dấu "Cần kiểm" bị LOẠI khỏi nguồn ghép')
    finally:
        os.chdir(cwd)

print(f'\n===== KẾT QUẢ: PASS {PASS} / FAIL {FAIL} =====')
sys.exit(1 if FAIL else 0)
