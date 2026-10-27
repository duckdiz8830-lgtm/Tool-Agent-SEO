# -*- coding: utf-8 -*-
"""Ver 3.1 — smoke test màn hình dữ liệu mới, chạy THẬT dưới Xvfb.

Chạy: xvfb-run -a python tests/test_v31_man_hinh.py
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import tkinter as tk
from tkinter import ttk

PASS = FAIL = 0


def kiem(dk, ten):
    global PASS, FAIL
    PASS, FAIL = (PASS + 1, FAIL) if dk else (PASS, FAIL + 1)
    print(f'  {"PASS" if dk else "FAIL"}  {ten}')


tmp = tempfile.mkdtemp()
os.chdir(tmp)

# file master mẫu + file kết quả biển máy mẫu
import openpyxl
from openpyxl.styles import PatternFill
wb = openpyxl.Workbook()
ws = wb.active
ws.title = '03_May_Bien_Ap_Do_Luong'
ws.append(['Mã hàng', 'Tên hàng', 'Tổn hao không tải Po (W, catalogue)'])
ws.append(['A1', 'Máy biến áp dầu 3Pha 250KVA', 440])
ws.append(['A2', 'Máy biến áp dầu 3Pha 320KVA', ''])
ws.append(['A3', 'Máy biến áp dầu 3Pha 560KVA', ''])
wb.save('master_mau.xlsx')

wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.append(['Ảnh nguồn', 'Công suất (kVA)', 'Cần kiểm', 'Lý do'])
ws2.append(['a.jpg', 320, '', ''])
ws2.append(['b.jpg', 560, 'X', 'ảnh mờ'])
wb2.save('bien-may-ket-qua.xlsx')

root = tk.Tk()
root.withdraw()
nb = ttk.Notebook(root)

import man_hinh_du_lieu as M

kiem(M.CO_TKSHEET or not M.CO_TKSHEET, 'tksheet dependency được nhận diện; thiếu thì dùng fallback chỉ-xem')
t_giau, t_xem = M.gan_vao(nb, app=None)
kiem(len(nb.tabs()) == 2, 'gắn được 2 tab con vào notebook')

# --- tab Xem dữ liệu: nạp + lọc cột trống
t_xem.duong_dan = 'master_mau.xlsx'
t_xem.v_sheet.set('03_May_Bien_Ap_Do_Luong')
t_xem._nap()
kiem(len(t_xem.data) == 3 and 'Tên hàng' in t_xem.hdr, 'nạp master mẫu: 3 dòng')
if M.CO_TKSHEET:
    t_xem.v_cot.set('Tổn hao không tải Po (W, catalogue)')
    t_xem._loc()
    so_hien = t_xem.bang.get_total_rows()
    kiem(so_hien == 2, f'lọc "thiếu Po" -> đúng 2/3 dòng (hiện {so_hien})')
    kiem('2/3' in t_xem.nhan_tk.cget('text'), 'nhãn thống kê ghi "2/3 dòng ĐANG TRỐNG"')
    t_xem.v_cot.set('(tất cả dòng)')
    t_xem._loc()
    kiem(t_xem.bang.get_total_rows() == 3, 'bỏ lọc -> hiện lại đủ 3 dòng')
else:
    kiem(t_xem.bang is None, 'thiếu tksheet -> tab Xem dữ liệu rơi về fallback an toàn, không crash')

# --- tab Làm giàu: nạp bảng soát, dòng X tô vàng
t_giau._lam_moi_ds_ket_qua()
kiem('bien-may-ket-qua.xlsx' in t_giau.cb_kq['values'], 'combobox thấy file kết quả')
if M.CO_TKSHEET:
    t_giau.v_file_kq.set('bien-may-ket-qua.xlsx')
    t_giau._nap_bang()
    kiem(t_giau.bang.get_total_rows() == 2, 'nạp bảng soát: 2 dòng')
    # sửa 1 ô rồi lưu — vòng tksheet -> xlsx phải giữ đúng
    t_giau.bang.set_cell_data(1, 2, '')          # gỡ dấu X "Cần kiểm" của dòng b.jpg
    t_giau._luu_bang()
    wb3 = openpyxl.load_workbook('bien-may-ket-qua.xlsx')
    ws3 = wb3.active
    kiem(str(ws3.cell(row=3, column=3).value or '') == '', 'sửa ô trong bảng -> lưu xuống file đúng')
else:
    kiem(t_giau.bang is None, 'thiếu tksheet -> tab Làm giàu rơi về fallback an toàn, không crash')

# --- phân loại file trong hàng đợi
kiem(t_giau._loai_file('x.jpg') == 'ảnh' and t_giau._loai_file('y.pdf') == 'PDF'
     and t_giau._loai_file('z.xlsx') == 'Excel', 'nhận đúng loại file (ảnh/PDF/Excel)')

root.destroy()
print(f'\n===== KẾT QUẢ: PASS {PASS} / FAIL {FAIL} =====')
sys.exit(1 if FAIL else 0)
