import os as _os, sys as _sys  # Ver 5.1 -- gọn thư mục: test chuyển vào tests/, thêm đường dẫn về thư mục gốc để import các module chính
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), '..'))
import os
import subprocess
import sys
import openpyxl

import nhap_phase2 as N
import chuan_hoa_master as CHM

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # Ver 5.1 -- test o tests/, du lieu/module chinh o thu muc goc (len 1 cap)
RA = '/tmp/test_v41_master_enriched.xlsx'

KQ = []


def test(ma, mota, ok):
    KQ.append((ma, mota, bool(ok)))
    print(f'[{"PASS" if ok else "FAIL"}] {ma}: {mota}')


if os.path.exists(RA):
    os.remove(RA)

# ============================================================ NHÓM A — chạy tool qua dòng lệnh thật (đúng cách Thịnh sẽ chạy)
kq = subprocess.run(
    [sys.executable, os.path.join(BASE, 'chuan_hoa_master.py'),
     '--nguon', os.path.join(BASE, 'du_lieu_goc'), '--ra', RA],
    capture_output=True, text=True)
test('CHM-A1', f'Chạy qua dòng lệnh (python3 chuan_hoa_master.py) không lỗi, exit={kq.returncode}',
     kq.returncode == 0 and os.path.exists(RA))

# ============================================================ NHÓM B — RÀNG BUỘC 1: sheet đầu tiên không được là Dashboard
wb = openpyxl.load_workbook(RA, read_only=True)
test('CHM-B1', f'Sheet đầu tiên là sheet dữ liệu thật, KHÔNG phải Dashboard: {wb.sheetnames[0]}',
     wb.sheetnames[0] == '01_Cap_Dien')
test('CHM-B2', 'nhan_dien_loai_nguon() nhận đúng định dạng 4_profile (không lẫn sang rmu_legacy)',
     N.nhan_dien_loai_nguon(RA) == '4_profile')

# ============================================================ NHÓM C — RÀNG BUỘC 2+3: header đúng dòng 1, thứ tự dòng + record_id giữ nguyên 100%
goc = N.import_4_profile(os.path.join(BASE, 'du_lieu_goc', 'MASTER_4_PROFILE_SAN_PHAM.xlsx'))
moi = N.import_4_profile(RA)
tat_ca_khop = True
for pid in goc:
    ids_goc = [r['record_id'] for r in goc[pid]['rows']]
    ids_moi = [r['record_id'] for r in moi[pid]['rows']]
    khop = ids_goc == ids_moi and len(ids_goc) == len(ids_moi)
    tat_ca_khop = tat_ca_khop and khop
    test(f'CHM-C-{pid}', f'record_id khớp 100% với file gốc ({len(ids_goc)} dòng, header không bị lệch dòng)', khop)
test('CHM-C-tong', 'TẤT CẢ 4 profile record_id khớp tuyệt đối (điều kiện sống còn để không mất liên kết bài đã đăng)',
     tat_ca_khop)

# ============================================================ NHÓM D — RÀNG BUỘC 4: không ghi đè giá trị cột gốc đã có
wb_goc = openpyxl.load_workbook(os.path.join(BASE, 'du_lieu_goc', '01_Cap_Dien.xlsx'), read_only=True, data_only=True)
rows_goc = list(wb_goc.active.iter_rows(values_only=True))
header_goc = rows_goc[0]
idx_goc = {h: i for i, h in enumerate(header_goc)}

wb_moi = openpyxl.load_workbook(RA, read_only=True, data_only=True)
rows_moi = list(wb_moi['01_Cap_Dien'].iter_rows(values_only=True))
header_moi = rows_moi[0]
idx_moi = {h: i for i, h in enumerate(header_moi)}

so_o_bi_sua = 0
for r_goc, r_moi in zip(rows_goc[1:], rows_moi[1:]):
    for c in ['Tiết Diện (mm²)', 'Cấp Điện Áp', 'Mã hàng', 'Tên hàng', 'Giá List (VNĐ)']:
        v_goc = r_goc[idx_goc[c]]
        if v_goc not in (None, '') and str(v_goc).strip() != str(r_moi[idx_moi[c]]).strip():
            so_o_bi_sua += 1
test('CHM-D1', f'KHÔNG có ô nào bị ghi đè sai (chỉ điền vào ô trống) — số ô sai: {so_o_bi_sua}', so_o_bi_sua == 0)

# ============================================================ NHÓM E — RÀNG BUỘC 5: regex có ranh giới đơn vị đúng (bug VA/KVA đã vá)
test('CHM-E1', 'VA (biến dòng) KHÔNG bị bắt nhầm thành V (điện áp)',
     CHM.RE_CAP_DIEN_AP.search('Biến dòng hạ thế 50/5A, 2.5VA W=1') is None)
m = CHM.RE_CAP_DIEN_AP.search('Máy biến áp dầu 1Pha 15KVA 12.7/0.23KV Loại Silic - TBD')
test('CHM-E2', f'KVA (công suất) không lấn vào kV (điện áp) — bắt đúng "12.7/0.23KV": {m.group(0) if m else None}',
     m is not None and '12.7/0.23' in m.group(0))

# ============================================================ NHÓM F — nội dung file: trạng thái/dashboard hợp lệ
ws_cd = wb_moi['01_Cap_Dien']
header = [c.value for c in next(ws_cd.iter_rows(min_row=1, max_row=1))]
idx = {h: i for i, h in enumerate(header)}
data = list(ws_cd.iter_rows(min_row=2, values_only=True))
so_confirmed = sum(1 for r in data if r[idx['Trạng Thái Kiểm Duyệt']] == 'CONFIRMED')
test('CHM-F1', f'Cáp điện có tỷ lệ CONFIRMED hợp lý (>50%, dữ liệu gốc vốn khá đầy đủ): {so_confirmed}/{len(data)}',
     so_confirmed / len(data) > 0.5)
test('CHM-F2', "Dashboard tồn tại và không phải sheet đầu", '00_Dashboard' in wb_moi.sheetnames
     and wb_moi.sheetnames.index('00_Dashboard') > 0)

print()
p = sum(1 for _, _, ok in KQ if ok)
print(f'TONG VER 4.1 CHUAN HOA MASTER: {p}/{len(KQ)} PASS')
if p != len(KQ):
    sys.exit(1)
