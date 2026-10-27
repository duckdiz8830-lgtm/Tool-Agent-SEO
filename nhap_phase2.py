# -*- coding: utf-8 -*-
"""
Phase 2 Lite — Bước 2: dispatcher import chung + data_checksum/dataset_id/
record_id (ARCHITECTURE-DRAFT-4-PROFILES.md mục 1, 3, 9).

Không sửa nhap_excel() cũ trong dulieu.py — gọi nguyên xi qua
import_rmu_legacy(). Đây là hàm mới song song, KHÔNG thay thế.
"""
import hashlib
import json
import re
import sqlite3
from pathlib import Path

import openpyxl

import profiles as P

BASE = Path(__file__).parent
DB_PHASE2 = BASE / 'phase2-du-lieu.db'

# Ver 1.7 -- Validation Gate: cung regex da verify that (163/165 dong dung)
# luc tu tay sua file CLEAN.xlsx truoc do -- tai dung o day de CANH BAO luc
# import, khong phai chi sua 1 lan roi thoi. Bat dang "X/YkV" hoac
# "X/Y kV" trong Ten hang (chap nhan ca dau phay thap phan kieu Viet).
_MAU_DIEN_AP_TRONG_TEN = re.compile(r'(\d+(?:[.,]\d+)?)\s*/\s*(\d+(?:[.,]\d+)?)\s*K[V]?\b',
                                     re.IGNORECASE)


def _so_tu_chuoi(v):
    if v is None:
        return None
    m = re.search(r'[\d.]+', str(v))
    return m.group(0) if m else None


def _quet_nghi_van_dien_ap(profile_id, rows):
    """Validation Gate -- CHI ap dung profile may bien ap (dung pham vi
    loi that da gap: Cap Dien Ap bi gan TRUNG so voi Cong Suat do loi
    enrich cu). Khong tu sua gi -- chi tra list nghi van de log canh bao
    + danh dau row['nghi_van_dien_ap']=True cho cho khac loc ra duyet tay.

    Logic: neu Ten hang co dang "X/YkV" ro rang, so voi Cap Dien Ap hien
    co -- KHONG khop (hoac Cap Dien Ap trung so voi Cong Suat -- dau hieu
    dac trung cua loi da gap) thi coi la nghi van."""
    if profile_id != 'may_bien_ap_do_luong':
        return []
    nghi_van = []
    for r in rows:
        ten = r.get('ten_goc', '') or ''
        cap_dien_ap = str(r.get('cap_dien_ap', '') or '')
        cong_suat_so = _so_tu_chuoi(r.get('cong_suat', ''))
        cap_dien_ap_so = _so_tu_chuoi(cap_dien_ap)
        m = _MAU_DIEN_AP_TRONG_TEN.search(ten)
        nghi_ngo = False
        de_xuat = None
        if cong_suat_so and cap_dien_ap_so and cong_suat_so == cap_dien_ap_so:
            nghi_ngo = True
            if m:
                de_xuat = f'{m.group(1).replace(",", ".")}/{m.group(2).replace(",", ".")} kV'
        elif m:
            a, b = m.group(1).replace(',', '.'), m.group(2).replace(',', '.')
            if a not in cap_dien_ap or b not in cap_dien_ap:
                nghi_ngo = True
                de_xuat = f'{a}/{b} kV'
        if nghi_ngo:
            r['nghi_van_dien_ap'] = True
            nghi_van.append({'ma_goc': r.get('ma_goc', ''), 'ten_goc': ten,
                             'cap_dien_ap_hien_tai': cap_dien_ap, 'de_xuat': de_xuat})
    return nghi_van

# Tên sheet trong MASTER hoặc tên file (bỏ đuôi .xlsx) -> profile_id.
# Ánh xạ TĨNH, không đoán bằng AI (đúng SPEC).
BAN_DO_PROFILE = {
    '01_cap_dien': 'cap_dien',
    '02_tu_va_thiet_bi_dong_cat': 'tu_va_thiet_bi_dong_cat',
    '03_may_bien_ap_do_luong': 'may_bien_ap_do_luong',
    '04_phu_kien_va_khac': 'phu_kien_va_khac',
}

# Tên sheet CHUẨN dùng để build record_id — cố định, không phụ thuộc file
# con đặt sheet là "Sheet1" hay MASTER đặt đúng tên. Nhờ vậy record_id của
# cùng 1 dòng dữ liệu logic GIỐNG HỆT NHAU dù import qua MASTER hay qua file
# con riêng — quan trọng để dedup/continuity không vỡ khi đổi cách nhập.
TEN_SHEET_CHUAN = {
    'cap_dien': '01_Cap_Dien',
    'tu_va_thiet_bi_dong_cat': '02_Tu_Va_Thiet_Bi_Dong_Cat',
    'may_bien_ap_do_luong': '03_May_Bien_Ap_Do_Luong',
    'phu_kien_va_khac': '04_Phu_Kien_Va_Khac',
}


def _slug_ten(s):
    return str(s).strip().lower().replace(' ', '_')


# ============================================================ nhận diện loại nguồn
def nhan_dien_loai_nguon(duong_dan):
    """Ver 1.6 — SỬA LỖI THẬT: bản cũ chỉ đọc sheet ĐẦU TIÊN, nên phụ thuộc
    thứ tự sheet trong file — 1 sheet phụ (VD nhật ký/log) lỡ nằm ở vị trí
    đầu là toàn bộ file bị nhận nhầm 'rmu_legacy', dù 4 sheet dữ liệu thật
    vẫn còn nguyên. Đã xảy ra thật với Thịnh 2 lần. Sửa: quét TẤT CẢ sheet,
    sheet NÀO có header 'Mã hàng'+'Tên hàng' thì coi là 4_profile — không
    còn phụ thuộc thứ tự sheet nữa.
    """
    wb = openpyxl.load_workbook(duong_dan, read_only=True, data_only=True)
    try:
        for ten_sheet in wb.sheetnames:
            hang_dau = next(wb[ten_sheet].iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not hang_dau:
                continue
            hang_dau_txt = {str(x).strip() for x in hang_dau if x is not None}
            if 'Mã hàng' in hang_dau_txt and 'Tên hàng' in hang_dau_txt:
                return '4_profile'
    finally:
        wb.close()
    # Không sheet nào khớp 4_profile -> coi là RMU cũ (fallback, không đoán thêm)
    return 'rmu_legacy'


def nhan_dien_profile_tu_ten(ten_sheet_hoac_file):
    key = _slug_ten(Path(ten_sheet_hoac_file).stem)
    if key in BAN_DO_PROFILE:
        return BAN_DO_PROFILE[key]
    return None


# ============================================================ checksum
def file_checksum(duong_dan):
    h = hashlib.sha256()
    with open(duong_dan, 'rb') as f:
        for chunk in iter(lambda: f.read(65536), b''):
            h.update(chunk)
    return 'sha256:' + h.hexdigest()


def data_checksum(rows_normalized, cot_import_keys):
    """Hash dữ liệu đã normalize, gồm cả thông số động.

    `cot_import_keys` chỉ xác định nhóm field cố định; JSON được sort key nên
    checksum không phụ thuộc thứ tự cột/thuộc tính. `thong_so_dong` cũng được
    canonical hóa để mọi thay đổi giá trị động đều đổi identity của dataset.
    """
    payload = []
    for row in rows_normalized:
        fixed = {k: row.get(k, '') for k in cot_import_keys}
        dynamic = row.get('thong_so_dong') or {}
        if not isinstance(dynamic, dict):
            dynamic = {}
        payload.append({
            'fixed': fixed,
            'dynamic': dict(sorted((str(k), v) for k, v in dynamic.items())),
        })
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(',', ':'))
    return 'sha256:' + hashlib.sha256(raw.encode('utf-8')).hexdigest()


# ============================================================ record_id
def record_id(profile_id, source_sheet, source_row, sku):
    return f'{profile_id}::{source_sheet}::{source_row}::{sku}'


def _doc_sheet_theo_header(ws, cot_import, cot_bat_buoc=None):
    """ws: openpyxl worksheet. cot_import: dict {tên cột Excel: field nội bộ}
    — TOÀN BỘ cột profile được PHÉP đọc nếu có mặt trong file.
    cot_bat_buoc: list tên cột Excel BẮT BUỘC phải có, thiếu là lỗi dừng
    ngay. Mặc định (None) = toàn bộ cot_import, giữ nguyên hành vi cũ cho
    profile chưa khai báo field optional (Ver 4.9.5 — 'phu_kien_va_khac' là
    profile đầu tiên tách 2 khái niệm này, xem PROFILES['phu_kien_va_khac']
    ['cot_bat_buoc'] trong profiles.py: 8 cột gốc bắt buộc, cột enrich
    catalogue KHÔNG bắt buộc — MASTER chưa enrich vẫn import được bình
    thường, chỉ là field đó trống).

    Trả list[(source_row, raw_dict)] — raw_dict giữ nguyên giá trị pandas/
    openpyxl thô, CHƯA normalize (normalize làm ở bước sau, tách riêng để dễ
    unit test từng bước độc lập).
    """
    hang = list(ws.iter_rows(values_only=True))
    if not hang:
        return []
    header = [str(x).strip() if x is not None else '' for x in hang[0]]
    idx = {ten: i for i, ten in enumerate(header)}
    bat_buoc = cot_bat_buoc if cot_bat_buoc is not None else list(cot_import)
    thieu = [ten for ten in bat_buoc if ten not in idx]
    if thieu:
        raise RuntimeError(f'Sheet "{ws.title}" thiếu cột bắt buộc: {thieu}')
    ra = []
    for r_idx, hang_du_lieu in enumerate(hang[1:], start=2):  # start=2: dòng 1 là header
        if hang_du_lieu is None or all(v is None for v in hang_du_lieu):
            continue
        raw = {ten: (hang_du_lieu[i] if i < len(hang_du_lieu) else None)
               for ten, i in idx.items()}
        if not raw.get('Mã hàng'):
            continue
        ra.append((r_idx, raw))
    return ra


# ============================================================ import 4-profile
def import_4_profile(duong_dan, chi_profile=None):
    """Đọc file MASTER (nhiều sheet) hoặc 1 file con (1 sheet).
    chi_profile: giới hạn chỉ đọc 1 profile_id (dùng khi test riêng lẻ file
    con, không bắt buộc khi đọc MASTER).

    Trả dict:
        {profile_id: {'rows': [...], 'source_sheet': ..., 'file_checksum': ...,
                       'data_checksum': ..., 'bao_cao': {...}}}
    Mỗi row trong 'rows' đã normalize đầy đủ + có record_id/source_row/source_sheet.
    """
    fc = file_checksum(duong_dan)
    wb = openpyxl.load_workbook(duong_dan, data_only=True)
    ra = {}
    for ten_sheet in wb.sheetnames:
        # Ưu tiên nhận diện theo tên SHEET (đúng cho MASTER). File con chỉ có
        # 1 sheet tên "Sheet1" không mang thông tin -> fallback theo tên FILE
        # nguồn (đúng nguyên tắc ARCHITECTURE mục 1.1: "profile lấy trực tiếp
        # từ tên sheet nguồn (MASTER) hoặc tên file nguồn (file con)").
        pid = nhan_dien_profile_tu_ten(ten_sheet) or nhan_dien_profile_tu_ten(duong_dan)
        if pid is None:
            continue
        if chi_profile and pid != chi_profile:
            continue
        cfg = P.PROFILES[pid]
        cot_import = cfg['cot_import']
        cot_import_keys = list(cot_import.values())
        cot_bat_buoc = cfg.get('cot_bat_buoc')  # None -> mặc định = toàn bộ cot_import
        thoat = _doc_sheet_theo_header(wb[ten_sheet], cot_import, cot_bat_buoc)
        ten_sheet_chuan = TEN_SHEET_CHUAN[pid]

        rows = []
        loi_thieu_ma = 0
        for source_row, raw in thoat:
            norm = P.normalize_row(raw, cot_import)
            if not norm.get('ma_goc'):
                loi_thieu_ma += 1
                continue
            norm['profile_id'] = pid
            norm['source_sheet'] = ten_sheet_chuan
            norm['source_row'] = source_row
            norm['record_id'] = record_id(pid, ten_sheet_chuan, source_row, norm['ma_goc'])
            norm['outline_id'] = P.xac_dinh_outline_id(pid, norm)
            norm['content_type'] = P.xac_dinh_content_type(pid, norm)
            rows.append(norm)

        dc = data_checksum(rows, cot_import_keys)

        # cảnh báo SKU trùng trong cùng profile — không gộp, chỉ cảnh báo
        thay_sku = {}
        sku_trung = []
        for r in rows:
            k = r['ma_goc']
            if k in thay_sku:
                sku_trung.append({'sku': k, 'record_id_1': thay_sku[k], 'record_id_2': r['record_id']})
            else:
                thay_sku[k] = r['record_id']

        # Ver 5.1 — VÁ: phát hiện khi audit lại tool (Audit-Checklist-Ver5.0
        # .xlsx mục Tool #3) — sku_trung ĐÃ được tính từ lâu (test_round2_
        # integration.py đã xác nhận đúng 17 cặp) nhưng CHỈ nằm trong báo
        # cáo, KHÔNG có hành động bảo vệ nào. Hậu quả thật: WooCommerce tra
        # sản phẩm theo SKU để update — 2 record khác nhau (VD BHSLBS24M có
        # SCADA/không SCADA, giá khác nhau 46tr/42tr) CÙNG SKU sẽ bị đăng
        # ĐÈ LÊN NHAU, mất 1 sản phẩm thật mà tool vẫn báo "THÀNH CÔNG".
        #
        # Vá: thêm field `publish_sku` — SKU DÙNG THẬT khi gửi lên
        # WooCommerce (ai_wp_adapter.py::dang_woo_that phải đọc field này,
        # KHÔNG đọc `ma_goc` thẳng cho tham số sku nữa). record_id đầu tiên
        # gặp mỗi mã giữ NGUYÊN sku gốc (không đổi hành vi cho ~99.7% SKU
        # không trùng); record thứ 2, 3... của CÙNG 1 mã trùng được gắn hậu
        # tố "-B2", "-B3"... `ma_goc` KHÔNG đổi — vẫn hiện đúng mã catalogue
        # thật trong content/bảng thông số, không đánh lừa khách hàng.
        dem_trung = {}
        for r in rows:
            k = r['ma_goc']
            dem_trung[k] = dem_trung.get(k, 0) + 1
            if dem_trung[k] == 1:
                r['publish_sku'] = k
            else:
                r['publish_sku'] = f'{k}-B{dem_trung[k]}'

        if sku_trung:
            print(f'  [CẢNH BÁO] {pid}: {len(sku_trung)} mã SKU bị trùng trong MASTER — '
                 f'đã gắn publish_sku riêng cho từng record để KHÔNG bị đè trên WooCommerce. '
                 f'Nên kiểm tra lại MASTER (Thịnh/Quỳnh) xem có phải lỗi nhập liệu không: '
                 f'{[s["sku"] for s in sku_trung[:5]]}{"..." if len(sku_trung) > 5 else ""}')

        # Ver 1.7 — VALIDATION GATE: phát hiện đúng dạng lỗi thật đã gặp
        # (165/245 SKU máy biến áp bị gán Cấp Điện Áp = Công Suất, do lỗi
        # enrich trước đó — xem CLEAN.xlsx/NHAT-KY-THAY-DOI.md Ver 5.4).
        # KHÔNG tự sửa, KHÔNG chặn import (giữ đúng triết lý "cảnh báo,
        # không âm thầm sửa" — giống hệt cách xử lý sku_trung ở trên) —
        # chỉ CẢNH BÁO RÕ + đánh dấu field `nghi_van_dien_ap=True` trên
        # từng row để chỗ khác (UI, báo cáo) có thể lọc ra duyệt tay.
        # Ver 2.4 — FAMILY ROUTER: kiểm tra chéo họ sản phẩm suy từ TÊN với
        # outline đang gán. Bắt lỗi SAI BẢN CHẤT (dịch vụ thí nghiệm bị
        # viết như bán thiết bị, chống sét van bị xếp vào tủ đóng cắt...).
        # CHỈ CẢNH BÁO, không tự sửa — giống cách xử lý sku_trung và
        # nghi_van_dien_ap: tool chỉ ra, người quyết định.
        try:
            import family_router as FR
            lech_ho = FR.quet_ca_dataset(rows)
        except Exception as e:
            lech_ho = []
            print(f'  [CẢNH BÁO] Không chạy được Family Router: {e}')
        if lech_ho:
            print(f'  [CẢNH BÁO] {pid}: {len(lech_ho)} SKU có dấu hiệu SAI HỌ SẢN PHẨM — '
                 f'tên hàng cho thấy loại khác với outline đang gán, bài sẽ viết sai '
                 f'bản chất. CHƯA tự sửa, cần Thịnh xem lại: '
                 f'{[l["ma_goc"] for l in lech_ho[:5]]}{"..." if len(lech_ho) > 5 else ""}')
            for l in lech_ho[:3]:
                print(f'      - {l["ma_goc"]}: "{l["ten_goc"][:50]}" '
                     f'-> đang {l["outline_hien_tai"]}, nên {l["outline_de_xuat"]}')

        nghi_van_dien_ap = _quet_nghi_van_dien_ap(pid, rows)
        if nghi_van_dien_ap:
            print(f'  [CẢNH BÁO] {pid}: {len(nghi_van_dien_ap)} dòng NGHI VẤN Cấp Điện Áp '
                 f'không khớp với Tên hàng (kiểu lỗi đã từng gặp: cột điện áp bị gán nhầm '
                 f'bằng số công suất) — ĐÃ IMPORT BÌNH THƯỜNG, chưa tự sửa gì, cần Thịnh tự '
                 f'soát lại: {[n["ma_goc"] for n in nghi_van_dien_ap[:5]]}'
                 f'{"..." if len(nghi_van_dien_ap) > 5 else ""}')

        ra[pid] = {
            'rows': rows,
            'source_sheet': ten_sheet_chuan,
            'file_checksum': fc,
            'data_checksum': dc,
            'bao_cao': {
                'tong_dong_doc_duoc': len(thoat),
                'hop_le': len(rows),
                'thieu_ma': loi_thieu_ma,
                'sku_trung': sku_trung,
                'nghi_van_dien_ap': nghi_van_dien_ap,
                'lech_ho_san_pham': lech_ho,
            },
        }
    wb.close()
    return ra


def import_rmu_legacy(duong_dan, **kwargs):
    """Đóng gói nguyên xi nhap_excel() cũ trong dulieu.py — KHÔNG sửa logic
    bên trong. Import trễ (lazy import) để module này không bắt buộc phải có
    sẵn dulieu.py lúc load (đỡ phụ thuộc vòng khi test riêng profiles.py)."""
    import dulieu as D
    return D.nhap_excel(duong_dan, **kwargs)


def dispatcher_import(duong_dan, **kwargs):
    """Điểm vào DUY NHẤT cho mọi việc nhập Excel — thay cho gọi thẳng
    D.nhap_excel() hay import_4_profile() tuỳ tiện ở UI."""
    loai = nhan_dien_loai_nguon(duong_dan)
    if loai == 'rmu_legacy':
        return {'loai': 'rmu_legacy', 'ket_qua': import_rmu_legacy(duong_dan, **kwargs)}
    elif loai == '4_profile':
        return {'loai': '4_profile', 'ket_qua': import_4_profile(duong_dan, **kwargs)}
    else:
        raise RuntimeError('Không nhận diện được cấu trúc file, không import liều.')


# ============================================================ R2-01: import orchestrator
def import_phase2_pipeline(file_path, confirm_new_dataset=False, chi_profile=None,
                           db_path=None, confirm_doi_cau_truc=False):
    """Hàm tổng DUY NHẤT cho luồng 4-profile — giao-dien.py chỉ gọi hàm này,
    không tự ghép logic import rời rạc (đúng R2-01).

    Thực hiện đủ 10 bước:
      1. Nhận diện MASTER hay file profile          (nhan_dien_loai_nguon)
      2. Đọc dữ liệu theo header                     (import_4_profile)
      3. normalize_cell                              (profiles.normalize_row, đã chạy trong bước 2)
      4. Gán profile_id/outline_id                   (profiles.xac_dinh_outline_id, đã chạy trong bước 2)
      5. Chạy validator + regex                      (validators.validate_rows)
      6. Tính file_checksum + data_checksum          (đã có trong kết quả bước 2)
      7. Xác định dataset_id                         (db_phase2.xac_nhan_dataset)
      8. Kiểm tra dataset cũ/mới                     (db_phase2.xac_nhan_dataset, có thể raise cảnh báo)
      9. Ghi database                                (db_phase2.ghi_san_pham, UPSERT)
      10. Trả báo cáo import, lỗi, cảnh báo

    Trả dict:
        {
          'ok': bool,
          'loai_nguon': '4_profile' | 'rmu_legacy',
          'ket_qua_theo_profile': {profile_id: {'dataset_id', 'trang_thai_dataset',
                                                  'so_dong', 'bao_cao_import', 'canh_bao'}},
          'loi': [str, ...],   # rỗng nếu không có lỗi chặn
        }

    Với nhánh rmu_legacy: chỉ chuyển tiếp qua import_rmu_legacy(), không đi
    qua profile/outline/validator/db_phase2 (nhánh đó dùng dulieu.py/du-lieu.db
    riêng, đúng thiết kế "legacy adapter", KHÔNG trộn 2 luồng).
    """
    import validators as V
    import db_phase2 as DB

    loai = nhan_dien_loai_nguon(file_path)
    bao_cao = {'ok': True, 'loai_nguon': loai, 'ket_qua_theo_profile': {}, 'loi': []}

    if loai == 'rmu_legacy':
        try:
            ket_qua_rmu = import_rmu_legacy(file_path)
            bao_cao['ket_qua_rmu_legacy'] = ket_qua_rmu
        except Exception as e:
            bao_cao['ok'] = False
            bao_cao['loi'].append(f'Import RMU legacy lỗi: {e}')
        return bao_cao

    if loai != '4_profile':
        bao_cao['ok'] = False
        bao_cao['loi'].append('Không nhận diện được cấu trúc file, không import liều.')
        return bao_cao

    # bước 2-6: đọc + normalize + gán profile/outline + tính checksum (import_4_profile làm hết)
    try:
        ket_qua = import_4_profile(file_path, chi_profile=chi_profile)
    except Exception as e:
        bao_cao['ok'] = False
        bao_cao['loi'].append(f'Đọc file lỗi: {e}')
        return bao_cao

    con = DB.ket_noi(db_path)
    try:
        for profile_id, d in ket_qua.items():
            muc = {'so_dong': len(d['rows']), 'bao_cao_import': d['bao_cao'], 'canh_bao': None}
            try:
                # bước 5: validator + regex
                rows = V.validate_rows(d['rows'])
                rows = V.chuan_hoa_dvt(rows)

                # bước 7-8: xác định dataset_id, có thể raise cảnh báo
                cot_import_keys = list(_cot_import_keys(profile_id))
                dataset_id, trang_thai_dataset = DB.xac_nhan_dataset(
                    con, profile_id, source_file=str(file_path),
                    file_checksum=d['file_checksum'], data_checksum=d['data_checksum'],
                    row_count=len(rows), column_structure=cot_import_keys,
                    confirm_new_dataset=confirm_new_dataset,
                    confirm_doi_cau_truc=confirm_doi_cau_truc)

                # bước 9: ghi database (UPSERT, không đụng bai_viet_phase2)
                DB.ghi_san_pham(con, dataset_id, rows)

                # bước 9b (Ver 4.4.1.7): dataset MỚI thì kéo bài đã đăng theo.
                # Không có bước này, bài cũ mồ côi ở dataset cũ và mọi chức
                # năng cần identity chết IM LẶNG — sự cố 13/08/2026.
                # Đặt SAU ghi_san_pham vì cần record_id của dataset mới để
                # đối chiếu. Không raise: import không được chết vì bước này.
                if trang_thai_dataset == 'moi_da_xac_nhan':
                    muc['chuyen_bai'] = DB.chuyen_bai_sang_dataset_moi(
                        con, profile_id, dataset_id)

                muc.update({'dataset_id': dataset_id, 'trang_thai_dataset': trang_thai_dataset,
                            'ok': True})
            except (DB.CanhBaoDatasetDoi, DB.CanhBaoCauTrucDoi) as e:
                cb = {'loai': type(e).__name__, 'thong_diep': str(e)}
                # Ver 4.5.1 — đưa cột thêm/bớt lên giao diện để người dùng
                # quyết định, thay vì chỉ ném một câu chung chung.
                cb['cot_them'] = list(getattr(e, 'them', []) or [])
                cb['cot_bot'] = list(getattr(e, 'bot', []) or [])
                muc.update({'ok': False, 'canh_bao': cb})
            except Exception as e:
                muc.update({'ok': False, 'canh_bao': {'loai': 'LoiKhac', 'thong_diep': str(e)}})
                bao_cao['loi'].append(f'{profile_id}: {e}')
            bao_cao['ket_qua_theo_profile'][profile_id] = muc
    finally:
        con.close()

    bao_cao['ok'] = bao_cao['ok'] and all(
        m.get('ok') for m in bao_cao['ket_qua_theo_profile'].values())
    return bao_cao


def _cot_import_keys(profile_id):
    import profiles as P
    return list(P.PROFILES[profile_id]['cot_import'].values())
