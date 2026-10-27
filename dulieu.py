# -*- coding: utf-8 -*-
"""
Nhập Excel bẩn, kho kiến thức theo dòng sản phẩm, quản ảnh, lưu trạng thái SQLite.
Không đụng tới phần gọi AI và đăng WordPress (nằm ở autopost.py).
"""
import json
import re
import sqlite3
import unicodedata
from pathlib import Path

BASE = Path(__file__).parent
DB = BASE / 'du-lieu.db'
KIEN_THUC = BASE / 'kienthuc'
THU_MUC_ANH = BASE / 'anh'

DUOI_ANH = ('.jpg', '.jpeg', '.png', '.webp', '.gif')

# tên sheet -> (hãng, series). Không khớp thì lấy nguyên tên sheet làm series.
BAN_DO_SHEET = {
    'schneider rm6': ('Schneider', 'RM6'),
    'schneider sm6': ('Schneider', 'SM6'),
    'abb safering safeplus': ('ABB', 'SafeRing / SafePlus'),
    'siemens 8djh': ('Siemens', '8DJH'),
    'taeyoung': ('TaeYoung', 'CD+'),
    'lucy aegis plus': ('Lucy', 'Aegis Plus'),
    'ormazabal': ('Ormazabal', 'CGM'),
}

# Chữ trong MÃ SẢN PHẨM tự nó nói lên series nào, ưu tiên hơn tên sheet chứa nó.
# Sếp dán nhầm 10 mã SM6 vào cả sheet RM6 lẫn sheet SM6 (file mẫu thật đã kiểm
# tra: dòng 42-51 sheet "Schneider RM6" = dòng 1-10 sheet "Schneider SM6", cùng
# mã, cùng cấu hình). Sheet quét trước không phải bằng chứng dòng đó thuộc dòng
# sản phẩm nào — chữ SM6/RM6 nằm ngay trong mã mới là bằng chứng đáng tin.
MA_TU_KHOA_SERIES = [('SM6', 'Schneider', 'SM6'), ('RM6', 'Schneider', 'RM6')]


def series_theo_ma(ma, hang, series):
    """Mã chứa từ khoá series rõ ràng thì ưu tiên từ khoá đó hơn tên sheet chứa nó."""
    ma_hoa = khong_dau(ma).upper()
    for tu_khoa, hang_dung, series_dung in MA_TU_KHOA_SERIES:
        if tu_khoa in ma_hoa and series != series_dung:
            return hang_dung, series_dung
    return hang, series


# ---------------------------------------------------------------- chuẩn hoá
def khong_dau(s):
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return s.replace('đ', 'd').replace('Đ', 'D')


def slugify(s, dai=90):
    s = re.sub(r'[^a-z0-9]+', '-', khong_dau(s).lower()).strip('-')
    return re.sub(r'-{2,}', '-', s)[:dai]


def chuan_don_vi(t):
    """24kv -> 24kV, 20KA/3s -> 20kA/3s, 630 a -> 630A, bỏ khoảng trắng thừa."""
    t = re.sub(r'\s+', ' ', str(t or '')).strip()
    t = re.sub(r'(\d(?:[.,]\d+)?)\s*k\s*v\b', r'\1kV', t, flags=re.I)
    t = re.sub(r'(\d+)\s*k\s*a\b', r'\1kA', t, flags=re.I)
    t = re.sub(r'(\d+)\s*a\b(?!\w)', r'\1A', t)
    t = re.sub(r'(\d+)\s*kva\b', r'\1kVA', t, flags=re.I)
    t = re.sub(r'kA\s*/\s*(\d*)\s*s\b', lambda m: f'kA/{m.group(1) or "1"}s', t, flags=re.I)
    return t


def khoa_trung(ma):
    """Khoá so trùng: bỏ dấu, bỏ khoảng trắng và ký tự phụ."""
    return re.sub(r'[^a-z0-9]', '', khong_dau(chuan_don_vi(ma)).lower())


# ---------------------------------------------------------------- bóc thông số
MAU_SO = {
    'cap_dien_ap': r'(\d{1,2}(?:[.,]\d)?\s*kV)',
    'dong_dinh_muc': r'(\d{3,4}\s*A)\b',
    'dong_ngan_mach': r'(\d{1,2}\s*kA\s*/\s*\d*\s*s)',
    'ro_le': r'\b(VIP\s?\d+|REJ\s?\d+|REF\s?\d+|SEPAM\s?\w+|VAMP\s?\d+|P3)\b',
}

LOAI_NGAN = [('cầu dao phụ tải', 'cầu dao phụ tải'), ('máy cắt', 'máy cắt'),
             ('đo lường', 'ngăn đo lường'), ('cầu chì', 'cầu chì'),
             ('cách ly', 'dao cách ly')]


def boc_thong_so(text):
    """Regex chạy trước, không tốn lượt AI. Tách thông số ra khỏi cột Cấu hình."""
    t = chuan_don_vi(text)
    ra = {}
    for ten, mau in MAU_SO.items():
        gt = []
        for x in re.findall(mau, t, re.I):
            x = chuan_don_vi(x).replace(' ', '')
            if x not in gt:
                gt.append(x)
        ra[ten] = ' / '.join(gt)

    # (?<![\d.,]) để "VIP400 Ngăn" không bị đọc thành "0 ngăn"
    m = re.search(r'(?<![\d.,])([1-9])\s*ng[aă]n', t, re.I)
    if m:
        ra['so_ngan'] = m.group(1)
    elif re.match(r'^\s*ng[aă]n\b', t, re.I) or re.search(r'\bng[aă]n t[uủ]\b', t, re.I):
        ra['so_ngan'] = '1'          # "Ngăn tủ máy cắt" là 1 ngăn đơn
    else:
        ra['so_ngan'] = ''

    low = t.lower()
    ra['loai_ngan'] = ', '.join(nhan for tu, nhan in LOAI_NGAN if tu in low)
    ra['cach_dien'] = ('SF6' if re.search(r'\bSF6\b', t, re.I)
                       else ('Chân không' if 'chân không' in low else ''))
    if re.search(r'không mở rộng', t, re.I):
        ra['mo_rong'] = 'Không'
    elif re.search(r'mở rộng', t, re.I):
        ra['mo_rong'] = 'Có'
    else:
        ra['mo_rong'] = ''
    return ra


def tach_model(ma):
    """Tách phần model ra khỏi chuỗi mã có lẫn rơ le và thông số."""
    t = chuan_don_vi(ma)
    t = re.sub(r'\bR[oơ]\s*le\b.*$', '', t, flags=re.I).strip()
    t = re.sub(r'\d{1,2}(?:[.,]\d)?kV.*$', '', t).strip()
    t = re.sub(r'[\-\+,]\s*$', '', t).strip()
    return t or chuan_don_vi(ma)


# ---------------------------------------------------------------- nhập Excel
COT_TRONG = ('gia_tham_khao', 'tinh_trang_kho', 'xuat_xu', 'bao_hanh', 'tieu_chuan',
             'kich_thuoc', 'trong_luong', 'anh_1_url', 'anh_1_caption', 'anh_2_url',
             'anh_2_caption', 'anh_3_url', 'anh_3_caption', 'catalogue_url',
             'ten_goi_khac', 'ung_dung', 'phu_kien_di_kem')


def goi_y_danh_muc(duong_dan):
    """Đoán mã danh mục từ tên file Excel. Sửa tay được ở giao diện."""
    return slugify(Path(duong_dan).stem) or 'chung'


def nhap_excel(duong_dan, sheets_chon=None, danh_muc=None):
    """Đọc file Excel nhiều sheet, trả (danh_sach_san_pham, bao_cao).

    danh_muc: ngăn kho theo ngành hàng. Không có thì lấy theo tên file.
    Thiếu cái này thì nạp file cáp điện đè lên file RMU là số liệu hai ngành
    trộn chung vào danh sách trắng chống bịa số.
    """
    danh_muc = danh_muc or goi_y_danh_muc(duong_dan)
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError(
            'Máy chưa có thư viện openpyxl nên không đọc được file Excel.\n\n'
            'Cách sửa: mở Command Prompt, gõ đúng dòng này rồi Enter:\n'
            '    python -m pip install openpyxl\n\n'
            'Xong thì mở lại phần mềm.')
    wb = openpyxl.load_workbook(duong_dan, data_only=True)
    bc = {'file': str(duong_dan), 'danh_muc': danh_muc, 'sheets': [], 'tong_dong': 0, 'hop_le': 0,
          'trung': [], 'ma_loi': [], 'sheet_rong': [], 'slug_da_sua': []}
    thay, ra = {}, []

    for sn in wb.sheetnames:
        if sheets_chon and sn not in sheets_chon:
            continue
        hang, series = BAN_DO_SHEET.get(khong_dau(sn).lower().strip(), ('', sn.strip()))
        dem = 0
        for r in wb[sn].iter_rows(values_only=True):
            # chỉ nhận dòng có cột đầu là số thứ tự
            if not r or r[0] is None:
                continue
            if not str(r[0]).replace('.0', '').strip().isdigit():
                continue
            dem += 1
            bc['tong_dong'] += 1
            ma = chuan_don_vi(str(r[1]) if len(r) > 1 and r[1] else '')
            cfg = chuan_don_vi(str(r[2]) if len(r) > 2 and r[2] else '')
            if not ma or ma in ('-', '_') or not re.search(r'[A-Za-z0-9]{2}', ma):
                bc['ma_loi'].append({'sheet': sn, 'ma': ma or '(rỗng)', 'cau_hinh': cfg[:60]})
                continue
            hang_dong, series_dong = series_theo_ma(ma, hang, series)
            k = khoa_trung(ma)
            trung = k in thay
            if trung:
                # KHÔNG bỏ dòng này — giữ làm record riêng theo yêu cầu, chỉ
                # cảnh báo trong báo cáo để biết chỗ mà kiểm tra tay nếu cần.
                bc['trung'].append({'ma': ma, 'sheet_moi': sn,
                                    'sheet_da_co': thay[k]['sheet_nguon']})

            ts = boc_thong_so(f'{ma} {cfg}')
            model = tach_model(ma)
            ngan = ts['so_ngan']
            # phần phân biệt: rơ le hoặc dòng ngắn mạch, để 2 mã cùng model không đè nhau
            khac = ts['ro_le'] or ts['dong_ngan_mach'] or ''
            if khac and khac.lower() in model.lower():
                khac = ''
            model_hien = re.sub(rf'^\s*{re.escape(series_dong)}[\s\-]*', '', model, flags=re.I) \
                or model
            ten = (f'Tủ RMU {hang_dong} {series_dong} {ngan} ngăn {ts["cap_dien_ap"] or ""} '
                   f'{model_hien} {khac}' if ngan
                   else f'{hang_dong} {series_dong} {model_hien} {khac}')
            ten = re.sub(r'\s{2,}', ' ', ten).strip()

            # sku là khoá lưu trữ + gốc để tạo URL. Mã trùng (đã có ở sheet khác)
            # thì thêm hậu tố tên sheet để không đè lên bản ghi trước — mỗi cái
            # vẫn ra một bài, một URL riêng. ma_goc giữ nguyên không đổi, dùng để
            # khớp ảnh và hiện trong bảng "Mã sản phẩm".
            sku = f'{ma} [{sn}]' if trung else ma

            sp = {
                # SKU lấy từ mã gốc đầy đủ, KHÔNG lấy model đã cắt, nếu không
                # RM6-NE-DIDI VIP400 và RM6-NE-DIDI VIP45 sẽ đè lên nhau
                'sku': sku, 'model': model, 'ma_goc': ma, 'ten_san_pham': ten,
                'hang': hang_dong, 'series': series_dong, 'dong_san_pham': series_dong,
                'cau_hinh_goc': cfg, 'sheet_nguon': sn, 'slug': slugify(ten),
                'danh_muc': danh_muc,
                **ts,
                **{c: '' for c in COT_TRONG},
            }
            thay[k] = sp
            ra.append(sp)
            bc['hop_le'] += 1
        bc['sheets'].append({'ten': sn, 'dong': dem, 'hang': hang, 'series': series})
        if dem == 0:
            bc['sheet_rong'].append(sn)

    # Slug phải duy nhất tuyệt đối, trùng là WordPress đẻ ra -2, -3.
    dung = {}
    for sp in ra:
        g = sp['slug']
        if g in dung:
            i = 2
            while f'{g}-{i}' in dung:
                i += 1
            sp['slug'] = f'{g}-{i}'
            bc['slug_da_sua'].append({'sku': sp['sku'], 'slug': sp['slug']})
        dung[sp['slug']] = sp['sku']
    return ra, bc


def tom_tat_bao_cao(bc):
    d = [f"Đọc {len(bc['sheets'])} sheet, {bc['tong_dong']} dòng.",
         f"  Hợp lệ                    : {bc['hop_le']}",
         f"  Trùng mã, giữ record riêng: {len(bc['trung'])}",
         f"  Mã không hợp lệ           : {len(bc['ma_loi'])}",
         f"  Sheet không dữ liệu       : {', '.join(bc['sheet_rong']) or 'không có'}",
         f"  Slug phải đánh số lại     : {len(bc['slug_da_sua'])}"]
    for t in bc['trung'][:12]:
        d.append(f"    trùng mã: {t['ma']}  ({t['sheet_da_co']} và {t['sheet_moi']}) "
                 f"-> vẫn tạo 2 record, tự kiểm tra xem có đúng ý không")
    if len(bc['trung']) > 12:
        d.append(f"    ... còn {len(bc['trung']) - 12} mã trùng nữa")
    for t in bc['ma_loi'][:6]:
        d.append(f"    mã lỗi: sheet {t['sheet']} -> {t['ma']}")
    return '\n'.join(d)


# ---------------------------------------------------------------- kho kiến thức
_CACHE_KT = {}


def _don_kien_thuc():
    """File .json kiến thức lỡ nằm ở thư mục gốc thì chuyển vào kienthuc/."""
    KIEN_THUC.mkdir(exist_ok=True)
    for f in BASE.glob('*.json'):
        try:
            d = json.loads(f.read_text(encoding='utf-8'))
        except Exception:
            continue
        if isinstance(d, dict) and 'series' in d and not (KIEN_THUC / f.name).exists():
            f.replace(KIEN_THUC / f.name)


def nap_kien_thuc(series=None, hang=None):
    """Đọc file kiến thức khớp series. CHỈ trả về các mục đã kiểm chứng."""
    _don_kien_thuc()
    key = slugify(f'{hang or ""} {series or ""}')
    if not key:
        return {}
    if key in _CACHE_KT:
        return _CACHE_KT[key]
    tot = {}
    for f in sorted(KIEN_THUC.glob('*.json')):
        try:
            d = json.loads(f.read_text(encoding='utf-8'))
        except Exception:
            continue
        nhan = slugify(f"{d.get('hang', '')} {d.get('series', '')}")
        chi_series = slugify(d.get('series', ''))
        if nhan and (nhan in key or key in nhan or (chi_series and chi_series in key)):
            tot = loc_da_kiem_chung(d)
            break
    _CACHE_KT[key] = tot
    return tot


def xoa_cache_kien_thuc():
    _CACHE_KT.clear()


def loc_da_kiem_chung(d):
    """Bỏ mọi mục da_kiem_chung = false. Chưa đối chiếu catalogue thì AI không được dùng."""
    if isinstance(d, dict):
        if 'gia_tri' in d and 'da_kiem_chung' in d:
            return d['gia_tri'] if d['da_kiem_chung'] else None
        ra = {}
        for k, v in d.items():
            v2 = loc_da_kiem_chung(v)
            if v2 not in (None, {}, []):
                ra[k] = v2
        return ra
    if isinstance(d, list):
        ra = [loc_da_kiem_chung(x) for x in d]
        return [x for x in ra if x not in (None, {}, [])]
    return d


def kien_thuc_thanh_chu(kt):
    """Biến kho kiến thức thành đoạn văn bản đưa vào prompt."""
    if not kt:
        return ''
    d = []
    if kt.get('tong_quan'):
        d.append(f"Tổng quan dòng sản phẩm: {kt['tong_quan']}")
    if kt.get('giai_thich_ma'):
        d.append('Ý nghĩa ký hiệu cấu hình: '
                 + '; '.join(f'{k} = {v}' for k, v in kt['giai_thich_ma'].items()))
    for ten, nhan in [('ung_dung', 'Ứng dụng điển hình'), ('uu_diem', 'Ưu điểm'),
                      ('phu_kien', 'Phụ kiện thường đi kèm'),
                      ('luu_y_chon', 'Lưu ý khi lựa chọn')]:
        if kt.get(ten):
            v = kt[ten]
            d.append(f'{nhan}: ' + ('; '.join(v) if isinstance(v, list) else str(v)))
    if kt.get('thong_so_chung'):
        d.append('Thông số chung của dòng (KHÔNG được dùng làm thông số riêng của model này): '
                 + '; '.join(f'{k} {v}' for k, v in kt['thong_so_chung'].items()))
    return '\n'.join(d)


def trang_thai_kien_thuc():
    """Liệt kê file kiến thức và mức độ đã kiểm chứng."""
    _don_kien_thuc()
    ra = []
    for f in sorted(KIEN_THUC.glob('*.json')):
        try:
            d = json.loads(f.read_text(encoding='utf-8'))
        except Exception as e:
            ra.append({'file': f.name, 'hang': 'LỖI JSON', 'series': str(e)[:50],
                       'tong_muc': 0, 'chua_kiem_chung': 0})
            continue
        tong = chua = 0

        def dem(x):
            nonlocal tong, chua
            if isinstance(x, dict):
                if 'da_kiem_chung' in x:
                    tong += 1
                    chua += 0 if x['da_kiem_chung'] else 1
                    return
                for v in x.values():
                    dem(v)
            elif isinstance(x, list):
                for v in x:
                    dem(v)
        dem(d)
        ra.append({'file': f.name, 'hang': d.get('hang', ''), 'series': d.get('series', ''),
                   'tong_muc': tong, 'chua_kiem_chung': chua})
    return ra


# ---------------------------------------------------------------- ảnh
def quet_anh():
    """Quét thư mục anh/, trả {khoa: [đường dẫn]}. Khoá lấy từ tên file."""
    THU_MUC_ANH.mkdir(exist_ok=True)
    ra = {}
    for f in THU_MUC_ANH.rglob('*'):
        if f.is_file() and f.suffix.lower() in DUOI_ANH:
            ra.setdefault(khoa_trung(f.stem), []).append(f)
    return ra


def gan_anh(sp, kho=None):
    """Tìm ảnh cho 1 sản phẩm: khớp mã trước, không có thì lùi về khớp dòng."""
    kho = quet_anh() if kho is None else kho
    for khoa in (khoa_trung(sp.get('sku', '')), khoa_trung(sp.get('ma_goc', '')),
                 khoa_trung(f"{sp.get('series', '')}{sp.get('so_ngan', '')}"),
                 khoa_trung(sp.get('series', ''))):
        if khoa and khoa in kho:
            return sorted(kho[khoa])[:3]
    return []


def bao_cao_anh(ds):
    kho = quet_anh()
    co = sum(1 for sp in ds if gan_anh(sp, kho) or (sp.get('anh_1_url') or '').strip())
    return {'tong_anh_trong_thu_muc': sum(len(v) for v in kho.values()),
            'san_pham_co_anh': co, 'san_pham_thieu_anh': len(ds) - co}


# ---------------------------------------------------------------- SQLite
def _cot(c, bang):
    return {r[1] for r in c.execute(f'PRAGMA table_info({bang})')}


def _co_bang(c, ten):
    return bool(c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name=?",
                          (ten,)).fetchone())


def _pk_cot(c, bang):
    """Tên các cột đang làm khoá chính, đúng thứ tự."""
    ds = [(r[5], r[1]) for r in c.execute(f'PRAGMA table_info({bang})') if r[5]]
    return [n for _, n in sorted(ds)]


def _nang_cap_db(c):
    """Nâng cấp DB cũ lên schema hiện tại, không mất lịch sử.

    V2      : bai_viet chưa có cột site  -> đổi tên bảng rồi chép sang bảng mới
    V2.1/2.3: san_pham khoá chính chỉ có 'sku' -> hai ngành hàng trùng SKU thì
              dòng sau đè dòng trước, mất im lặng. Xây lại bảng với khoá kép
              (danh_muc, sku).
    """
    if _co_bang(c, 'bai_viet') and 'site' not in _cot(c, 'bai_viet'):
        c.executescript('ALTER TABLE bai_viet RENAME TO bai_viet_cu;')
        c.commit()

    if _co_bang(c, 'san_pham') and _pk_cot(c, 'san_pham') != ['danh_muc', 'sku']:
        cot_hien = _cot(c, 'san_pham')
        dm_bt = 'danh_muc' if 'danh_muc' in cot_hien else "'chung'"
        chon_bt = 'chon' if 'chon' in cot_hien else '1'
        ngay_bt = 'ngay_nhap' if 'ngay_nhap' in cot_hien else 'CURRENT_TIMESTAMP'
        c.executescript('ALTER TABLE san_pham RENAME TO san_pham_cu;')
        c.executescript('''
        CREATE TABLE san_pham (
            danh_muc TEXT NOT NULL DEFAULT 'chung', sku TEXT NOT NULL, du_lieu TEXT,
            chon INTEGER DEFAULT 1, ngay_nhap TEXT DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (danh_muc, sku));
        ''')
        c.execute(f'''INSERT OR IGNORE INTO san_pham(danh_muc, sku, du_lieu, chon, ngay_nhap)
                      SELECT COALESCE({dm_bt}, 'chung'), sku, du_lieu,
                             COALESCE({chon_bt}, 1), {ngay_bt} FROM san_pham_cu''')
        c.executescript('DROP TABLE san_pham_cu;')
        c.commit()


# các chặng theo dõi trên dashboard, thêm cột mới thì khai ở đây
COT_CHANG = [('kiem_tra', "TEXT DEFAULT ''"), ('wp_trang_thai', "TEXT DEFAULT ''"),
             ('so_tu', 'INTEGER DEFAULT 0'), ('so_anh', 'INTEGER DEFAULT 0'),
             ('du_lieu_day', "TEXT DEFAULT ''"), ('danh_muc', "TEXT DEFAULT ''"),
             ('meta_title', "TEXT DEFAULT ''"), ('meta_description', "TEXT DEFAULT ''"),
             ('van_bien_tap', "TEXT DEFAULT ''")]


def mo_db():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    _nang_cap_db(c)
    c.executescript('''
    CREATE TABLE IF NOT EXISTS san_pham (
        danh_muc TEXT NOT NULL DEFAULT 'chung', sku TEXT NOT NULL, du_lieu TEXT,
        chon INTEGER DEFAULT 1, ngay_nhap TEXT DEFAULT CURRENT_TIMESTAMP,
        PRIMARY KEY (danh_muc, sku));
    CREATE TABLE IF NOT EXISTS bai_viet (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        site TEXT DEFAULT 'site1', sku TEXT, mau TEXT, slug TEXT, tieu_de TEXT, html TEXT,
        trang_thai TEXT, ghi_chu TEXT, wp_id INTEGER, wp_url TEXT,
        ngay TEXT DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(site, sku, mau));
    CREATE INDEX IF NOT EXISTS idx_slug ON bai_viet(slug);
    CREATE INDEX IF NOT EXISTS idx_dm ON san_pham(danh_muc);
    CREATE TABLE IF NOT EXISTS anh_wp (
        duong_dan TEXT, site TEXT, wp_id INTEGER, url TEXT, hash TEXT DEFAULT '',
        PRIMARY KEY (duong_dan, site));
    ''')
    if 'hash' not in _cot(c, 'anh_wp'):
        c.execute("ALTER TABLE anh_wp ADD COLUMN hash TEXT DEFAULT ''")
    thieu = _cot(c, 'bai_viet')
    for ten, kieu in COT_CHANG:
        if ten not in thieu:
            c.execute(f'ALTER TABLE bai_viet ADD COLUMN {ten} {kieu}')
    if _co_bang(c, 'bai_viet_cu'):
        c.executescript('''
        INSERT OR IGNORE INTO bai_viet
            (site,sku,mau,slug,tieu_de,html,trang_thai,ghi_chu,wp_id,wp_url,ngay)
        SELECT 'site1',sku,mau,slug,tieu_de,html,trang_thai,ghi_chu,wp_id,wp_url,ngay
        FROM bai_viet_cu;
        DROP TABLE bai_viet_cu;
        ''')
    c.commit()
    return c


def luu_san_pham(ds, danh_muc=None):
    c = mo_db()
    for sp in ds:
        dm = sp.get('danh_muc') or danh_muc or 'chung'
        sp['danh_muc'] = dm
        # nhập lại thì giữ lựa chọn cũ nếu (danh_muc, SKU) đã tồn tại, mới thì mặc định chọn
        cu = c.execute('SELECT chon FROM san_pham WHERE danh_muc=? AND sku=?',
                       (dm, sp['sku'])).fetchone()
        chon = cu['chon'] if cu is not None else 1
        c.execute('INSERT OR REPLACE INTO san_pham(sku, du_lieu, danh_muc, chon) VALUES (?,?,?,?)',
                  (sp['sku'], json.dumps(sp, ensure_ascii=False), dm, chon))
    c.commit()
    n = c.execute('SELECT COUNT(*) FROM san_pham').fetchone()[0]
    c.close()
    return n


def doc_san_pham_db(danh_muc=None, chi_lay_da_chon=False):
    """danh_muc: None hoặc [] = lấy tất cả. Chuỗi hoặc danh sách = lọc theo ngành hàng.

    chi_lay_da_chon=True: chỉ trả sản phẩm đang bấm chọn ở tab Sản phẩm. Mỗi dict
    có thêm khoá 'chon' (1/0) để giao diện tô đúng trạng thái.
    """
    if not DB.exists():
        return []
    c = mo_db()
    dk, tham = [], []
    if danh_muc:
        dm = [danh_muc] if isinstance(danh_muc, str) else list(danh_muc)
        dk.append(f"danh_muc IN ({','.join('?' * len(dm))})")
        tham += dm
    if chi_lay_da_chon:
        dk.append('chon=1')
    q = 'SELECT du_lieu, chon FROM san_pham'
    if dk:
        q += ' WHERE ' + ' AND '.join(dk)
    ra = []
    for r in c.execute(q, tham):
        sp = json.loads(r['du_lieu'])
        sp['chon'] = r['chon']
        ra.append(sp)
    c.close()
    return ra


def dat_chon(khoa_list, chon):
    """Bật/tắt cờ 'cần viết'. khoa_list: [(danh_muc, sku), ...]. chon: True/False.

    Khoá theo cặp (danh_muc, sku), không phải sku đơn — hai ngành hàng trùng SKU
    thì bật/tắt một cái không được đụng tới cái kia.
    """
    if not khoa_list:
        return
    c = mo_db()
    c.executemany('UPDATE san_pham SET chon=? WHERE danh_muc=? AND sku=?',
                  [(1 if chon else 0, dm, sku) for dm, sku in khoa_list])
    c.commit()
    c.close()


def dao_chon(danh_muc, sku):
    c = mo_db()
    c.execute('UPDATE san_pham SET chon = 1 - chon WHERE danh_muc=? AND sku=?', (danh_muc, sku))
    c.commit()
    r = c.execute('SELECT chon FROM san_pham WHERE danh_muc=? AND sku=?',
                  (danh_muc, sku)).fetchone()
    c.close()
    return r['chon'] if r else None


def cac_danh_muc():
    """[(mã danh mục, số sản phẩm)] để đổ vào ô chọn trên giao diện."""
    if not DB.exists():
        return []
    c = mo_db()
    ra = [(r[0] or 'chung', r[1]) for r in c.execute(
        'SELECT danh_muc, COUNT(*) FROM san_pham GROUP BY danh_muc ORDER BY 2 DESC')]
    c.close()
    return ra


def ghi_bai(site, sku, mau, slug, tieu_de, html, trang_thai, ghi_chu='',
            wp_id=None, wp_url='', **chang):
    """chang: kiem_tra, wp_trang_thai, so_tu, so_anh, du_lieu_day, danh_muc."""
    ten_chang = [t for t, _ in COT_CHANG]
    gt = [chang.get(t, '' if t not in ('so_tu', 'so_anh') else 0) for t in ten_chang]
    dat = ', '.join(f'{t}=excluded.{t}' for t in ten_chang)
    c = mo_db()
    c.execute(f'''INSERT INTO bai_viet
                   (site,sku,mau,slug,tieu_de,html,trang_thai,ghi_chu,wp_id,wp_url,
                    {",".join(ten_chang)})
                 VALUES (?,?,?,?,?,?,?,?,?,?,{",".join("?" * len(ten_chang))})
                 ON CONFLICT(site,sku,mau) DO UPDATE SET
                   slug=excluded.slug, tieu_de=excluded.tieu_de, html=excluded.html,
                   trang_thai=excluded.trang_thai, ghi_chu=excluded.ghi_chu,
                   wp_id=COALESCE(excluded.wp_id, bai_viet.wp_id),
                   wp_url=CASE WHEN excluded.wp_url<>'' THEN excluded.wp_url
                               ELSE bai_viet.wp_url END,
                   {dat}, ngay=CURRENT_TIMESTAMP''',
              [site, sku, mau, slug, tieu_de, html, trang_thai, ghi_chu, wp_id, wp_url] + gt)
    c.commit()
    c.close()


def doc_bai_viet(site=None, danh_muc=None):
    """Toàn bộ bài đã có trong kho, để dựng dashboard và xuất Excel."""
    if not DB.exists():
        return []
    c = mo_db()
    dk, tham = [], []
    if site:
        dk.append('site=?')
        tham.append(site)
    if danh_muc:
        dk.append('danh_muc=?')
        tham.append(danh_muc)
    q = ('SELECT site,sku,mau,slug,tieu_de,trang_thai,kiem_tra,wp_trang_thai,so_tu,so_anh,'
         'du_lieu_day,danh_muc,ghi_chu,wp_url,wp_id,ngay,meta_title,meta_description,'
         'LENGTH(html) AS so_ky_tu FROM bai_viet')
    if dk:
        q += ' WHERE ' + ' AND '.join(dk)
    ra = [dict(r) for r in c.execute(q + ' ORDER BY ngay DESC', tham)]
    c.close()
    return ra


def slug_trung():
    """Slug bị dùng cho nhiều SKU khác nhau trong cùng site. Phải luôn rỗng."""
    if not DB.exists():
        return []
    c = mo_db()
    ra = [dict(r) for r in c.execute(
        'SELECT site, slug, COUNT(DISTINCT sku) n, GROUP_CONCAT(DISTINCT sku) ds '
        'FROM bai_viet GROUP BY site, slug HAVING n > 1')]
    c.close()
    return ra


# ---------------------------------------------------------------- vân tay thông số
# Hai sản phẩm giống nhau ở đúng các cột này thì rất có thể là cùng một sản phẩm
# hoặc dữ liệu đang trống, không phải "tự nhiên trùng". Chỉ để CẢNH BÁO, không chặn.
COT_VAN_TAY = ('cap_dien_ap', 'dong_dinh_muc', 'dong_ngan_mach', 'ro_le',
               'so_ngan', 'loai_ngan', 'cach_dien', 'mo_rong')


def van_tay_thong_so(row):
    """Băm 8 cột thông số quan trọng thành 1 chuỗi. Rỗng hết thì trả '' (không
    tính vân tay cho dòng không có thông số nào, tránh báo trùng oan hàng loạt)."""
    gt = [str(row.get(k, '') or '').strip().lower() for k in COT_VAN_TAY]
    if not any(gt):
        return ''
    return '|'.join(gt)


def nhom_trung_van_tay(danh_sach_san_pham):
    """{sku: [sku_khac_cung_van_tay, ...]} — chỉ gồm SKU thật sự có vân tay trùng."""
    theo_vt = {}
    for r in danh_sach_san_pham:
        vt = van_tay_thong_so(r)
        if vt:
            theo_vt.setdefault(vt, []).append(r['sku'])
    ra = {}
    for sku_list in theo_vt.values():
        if len(sku_list) > 1:
            for sku in sku_list:
                ra[sku] = [s for s in sku_list if s != sku]
    return ra


# ---------------------------------------------------------------- văn biên tập cũ
def van_ban_cac_bai(site, tru_sku=None):
    """[(sku, van_bien_tap)] các bài đã sinh trong site này, để so trùng văn bản.
    Bỏ qua bài của chính sku đang sinh lại (tru_sku), và bài chưa có văn bản lưu
    (bản cũ trước khi có cột này, hoặc bài bị chặn chưa từng ra văn)."""
    if not DB.exists():
        return []
    c = mo_db()
    dk, tham = ["site=?", "van_bien_tap<>''"], [site]
    if tru_sku:
        dk.append('sku<>?')
        tham.append(tru_sku)
    ra = [(r['sku'], r['van_bien_tap']) for r in c.execute(
        f'SELECT sku, van_bien_tap FROM bai_viet WHERE {" AND ".join(dk)}', tham)]
    c.close()
    return ra


def da_lam(site, mau):
    """SKU đã xong với site + mẫu này, dùng để chạy tiếp sau khi tắt máy."""
    c = mo_db()
    ra = {r['sku'] for r in c.execute(
        "SELECT sku FROM bai_viet WHERE site=? AND mau=? "
        "AND trang_thai IN ('DA_SINH','DA_DANG')", (site, mau))}
    c.close()
    return ra


def wp_id_da_luu(site, sku, mau):
    """wp_id đã lưu từ lần đăng trước, nếu có. Đáng tin hơn hỏi WordPress theo
    slug lúc đang đăng, vì không phụ thuộc mạng ngay lúc này — tránh đẻ bài -2
    khi GET kiểm tra slug bị lỗi mạng ngay giữa lúc đang thử đăng."""
    c = mo_db()
    r = c.execute('SELECT wp_id FROM bai_viet WHERE site=? AND sku=? AND mau=? '
                  'AND wp_id IS NOT NULL', (site, sku, mau)).fetchone()
    c.close()
    return r['wp_id'] if r else None


def nho_anh_wp(duong_dan, site, file_hash=None):
    """Ảnh đã tải lên site này chưa? Nếu có hash thì bytes phải khớp."""
    c = mo_db()
    if file_hash is None:
        r = c.execute('SELECT wp_id, url FROM anh_wp WHERE duong_dan=? AND site=?',
                      (str(duong_dan), site)).fetchone()
    else:
        r = c.execute('SELECT wp_id, url FROM anh_wp WHERE duong_dan=? AND site=? AND hash=?',
                      (str(duong_dan), site, str(file_hash))).fetchone()
    c.close()
    return (r['wp_id'], r['url']) if r else None


def luu_anh_wp(duong_dan, site, wp_id, url, file_hash=''):
    c = mo_db()
    c.execute('INSERT OR REPLACE INTO anh_wp(duong_dan,site,wp_id,url,hash) VALUES (?,?,?,?,?)',
              (str(duong_dan), site, wp_id, url, str(file_hash or '')))
    c.commit()
    c.close()


def xoa_cache_anh_wp_theo_media(site, wp_id):
    """4.4.1.3 — xoá cache upload khi WordPress báo media ID đã stale.

    Chỉ xoá đúng site + media ID; không đụng cache site khác và không xoá
    file local. Lần sau uploader sẽ POST lại asset thay vì trả ID chết.
    """
    c = mo_db()
    try:
        cur = c.execute('DELETE FROM anh_wp WHERE site=? AND CAST(wp_id AS TEXT)=?',
                        (site, str(wp_id)))
        c.commit()
        return int(cur.rowcount or 0)
    finally:
        c.close()


def thong_ke(site=None):
    c = mo_db()
    d = {'san_pham': c.execute('SELECT COUNT(*) FROM san_pham').fetchone()[0]}
    sql = 'SELECT trang_thai, COUNT(*) n FROM bai_viet'
    tham = ()
    if site:
        sql += ' WHERE site=?'
        tham = (site,)
    for r in c.execute(sql + ' GROUP BY trang_thai', tham):
        d[r['trang_thai']] = r['n']
    c.close()
    return d
