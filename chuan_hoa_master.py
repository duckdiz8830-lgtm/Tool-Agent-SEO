# -*- coding: utf-8 -*-
"""
Ver 4.1 — CHUẨN HOÁ MASTER: đọc 4 file `du_lieu_goc/0X_*.xlsx`, làm giàu
thêm cột thông số kỹ thuật trích được bằng rule cố định (KHÔNG suy đoán,
KHÔNG gọi AI), gắn cờ CONFIRMED/CAN_KIEM + ghi chú, đánh dấu (không tự sửa)
nghi ngờ phân loại sai và trùng mã/tên, thêm Dashboard tổng quan.

RÀNG BUỘC BẤT BIẾN (đã vỡ 1 lần, GPT review bắt được — xem CHANGELOG-VER41):
  1. Sheet đầu tiên trong file xuất ra PHẢI là sheet dữ liệu thật (KHÔNG
     phải Dashboard) — nhap_phase2.nhan_dien_loai_nguon() chỉ đọc
     wb.sheetnames[0] để nhận diện định dạng file.
  2. Header PHẢI nằm ở DÒNG 1 của mỗi sheet dữ liệu, KHÔNG được chèn dòng
     tiêu đề phía trên — nhap_phase2._doc_sheet_theo_header() giả định cứng
     header ở hang[0], không tìm kiếm động.
  3. Thứ tự dòng dữ liệu trong mỗi sheet PHẢI giữ NGUYÊN 100% so với file
     gốc — record_id = f'{profile_id}::{sheet}::{source_row}::{sku}' phụ
     thuộc trực tiếp vào vị trí dòng. Đổi thứ tự = mất liên kết với mọi
     bài đã đăng WordPress thật.
  4. KHÔNG được ghi đè giá trị đã có sẵn ở cột gốc — chỉ điền vào ô trống.
  5. Regex trích số PHẢI có ranh giới đơn vị đúng (\\b + lookahead khi cần)
     — bài học thật: "2.5VA" (công suất biểu kiến biến dòng) từng bị bắt
     nhầm thành "2.5V" (điện áp) vì thiếu ranh giới, dẫn tới field kỹ thuật
     SAI vẫn bị đóng dấu CONFIRMED. Xem test_v41_chuan_hoa_master.py.

Dùng dòng lệnh:
    python3 chuan_hoa_master.py --nguon du_lieu_goc --ra MASTER_ENRICHED.xlsx
"""
import argparse
import collections
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

F = 'Arial'
NAVY = '1F3864'
WHITE = 'FFFFFF'
LBLUE = 'DCE6F1'
GREEN = 'C6EFCE'
YELLOW = 'FFEB9C'
RED = 'FFC7CE'
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

TEN_FILE = {
    'cap_dien': ('01_Cap_Dien.xlsx', '01_Cap_Dien', 'Cáp Điện'),
    'tu_va_thiet_bi_dong_cat': ('02_Tu_Va_Thiet_Bi_Dong_Cat.xlsx', '02_Tu_Va_Thiet_Bi_Dong_Cat', 'Tủ Và Thiết Bị Đóng Cắt'),
    'may_bien_ap_do_luong': ('03_May_Bien_Ap_Do_Luong.xlsx', '03_May_Bien_Ap_Do_Luong', 'Máy Biến Áp Đo Lường'),
    'phu_kien_va_khac': ('04_Phu_Kien_Va_Khac.xlsx', '04_Phu_Kien_Va_Khac', 'Phụ Kiện Và Khác'),
}

# regex trích xuất — xem mục 5 ràng buộc ở trên, đã kiểm bằng test thật
RE_TIET_DIEN = re.compile(r'(\d+(?:[.,]\d+)?(?:\s*[xX]\s*\d+(?:[.,]\d+)?){0,2})\s*mm', re.IGNORECASE)
RE_CAP_DIEN_AP = re.compile(r'\(?\s*(\d+(?:[.,]\d+)?(?:\s*/\s*\d+(?:[.,]\d+)?)?)\s*(kV|V)\b(?!A)\s*\)?', re.IGNORECASE)
RE_TIEU_CHUAN = re.compile(r'\b(TCVN|IEC|JIS|BS)[\s-]?\d+(?:[-.:]\d+)*', re.IGNORECASE)
RE_DONG_DINH_MUC = re.compile(r'(\d+(?:[.,]\d+)?)\s*A\b')
RE_CONG_SUAT = re.compile(r'(\d+(?:[.,]\d+)?)\s*KVA', re.IGNORECASE)
RE_LOAI_MBA = re.compile(r'\b(Silic|Amorphous)\b', re.IGNORECASE)

TU_KHOA_NGHI_NGO = {
    'tu_va_thiet_bi_dong_cat': [
        ('elbow', 'Đầu Elbow là phụ kiện nối cáp, không phải thiết bị đóng cắt'),
        ('đầu cáp', 'Đầu cáp là phụ kiện nối cáp, không phải thiết bị đóng cắt'),
        ('cosse', 'Cosse là phụ kiện đấu nối, không phải thiết bị đóng cắt'),
    ],
    'phu_kien_va_khac': [
        ('máy biến áp', 'Tên nhắc "máy biến áp" nhưng đang ở nhóm phụ kiện'),
        ('mba ', 'Tên nhắc MBA nhưng đang ở nhóm phụ kiện'),
    ],
}


def s(v):
    return None if v is None else str(v).strip()


def doc_sheet(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    return list(rows[0]), [list(r) for r in rows[1:]]


def _header_row(ws, r, values):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=F, size=9, bold=True, color=NAVY)
        c.fill = PatternFill('solid', fgColor=LBLUE)
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = border
    ws.row_dimensions[r].height = 32


def doc_sheet_tu_workbook(wb, ten_sheet):
    """Ver 4.8 — đọc 1 sheet từ workbook ĐÃ MỞ SẴN (dùng khi nguồn là 1 file
    gộp nhiều sheet, không phải 4 file riêng — tránh mở lại file nhiều lần)."""
    ws = wb[ten_sheet]
    rows = list(ws.iter_rows(values_only=True))
    return list(rows[0]), [list(r) for r in rows[1:]]


def _nap_du_lieu_nguon(nguon):
    """Ver 4.8 — nhận CẢ 2 kiểu nguồn, không bắt Thịnh phải nhớ đúng kiểu:
      (1) Thư mục chứa 4 file riêng (du_lieu_goc/01_Cap_Dien.xlsx...) — kiểu cũ.
      (2) 1 file .xlsx GỘP sẵn 4 sheet (MASTER_4_PROFILE_SAN_PHAM.xlsx, đúng
          file Thịnh thực tế đang dùng hằng ngày) — kiểu mới.
    Tự nhận diện bằng cách kiểm tra `nguon` là thư mục hay file."""
    nguon = Path(nguon)
    du_lieu = {}
    if nguon.is_dir():
        for pid, (ten_file, sheet_chuan, nhan) in TEN_FILE.items():
            duong_dan = nguon / ten_file
            if not duong_dan.exists():
                raise FileNotFoundError(f'Không tìm thấy {duong_dan} — thư mục nguồn phải chứa đủ 4 file '
                                        f'{", ".join(t[0] for t in TEN_FILE.values())}.')
            header, data = doc_sheet(duong_dan)
            du_lieu[pid] = {'header': header, 'data': data, 'sheet_chuan': sheet_chuan, 'nhan': nhan}
    elif nguon.is_file():
        wb = openpyxl.load_workbook(nguon, read_only=True, data_only=True)
        thieu_sheet = [sc for _, (_, sc, _) in TEN_FILE.items() if sc not in wb.sheetnames]
        if thieu_sheet:
            raise ValueError(f'File "{nguon.name}" thiếu sheet: {", ".join(thieu_sheet)}. '
                             f'File nguồn phải có đủ 4 sheet đúng tên: '
                             f'{", ".join(t[1] for t in TEN_FILE.values())}.')
        for pid, (ten_file, sheet_chuan, nhan) in TEN_FILE.items():
            header, data = doc_sheet_tu_workbook(wb, sheet_chuan)
            du_lieu[pid] = {'header': header, 'data': data, 'sheet_chuan': sheet_chuan, 'nhan': nhan}
    else:
        raise FileNotFoundError(f'Không tìm thấy nguồn: {nguon} (không phải thư mục cũng không phải file).')
    return du_lieu


def chuan_hoa(nguon, duong_dan_ra):
    """`nguon`: thư mục chứa 4 file riêng HOẶC 1 file .xlsx gộp sẵn 4 sheet
    — xem `_nap_du_lieu_nguon()`. Tên tham số đổi từ `thu_muc_nguon` thành
    `nguon` (Ver 4.8) để phản ánh đúng việc không còn bắt buộc là thư mục;
    lời gọi cũ dùng keyword `thu_muc_nguon=...` sẽ lỗi — dùng positional
    hoặc `nguon=...`."""
    du_lieu = _nap_du_lieu_nguon(nguon)

    dem_ma, dem_ten = collections.Counter(), collections.Counter()
    for pid, d in du_lieu.items():
        idx = {h: i for i, h in enumerate(d['header'])}
        for r in d['data']:
            ma, ten = s(r[idx['Mã hàng']]), s(r[idx['Tên hàng']])
            if ma:
                dem_ma[ma] += 1
            if ten:
                dem_ten[ten] += 1

    wb_moi = openpyxl.Workbook()
    wb_moi.remove(wb_moi.active)

    for pid, (ten_file, sheet_chuan, nhan) in TEN_FILE.items():
        header = du_lieu[pid]['header']
        data = du_lieu[pid]['data']
        idx = {h: i for i, h in enumerate(header)}

        cot_moi = ['Tiêu Chuẩn (trích từ tên)']
        if pid == 'tu_va_thiet_bi_dong_cat':
            cot_moi = ['Cấp Điện Áp (trích từ tên)', 'Dòng Định Mức A (trích từ tên)',
                      'Tiết Diện liên quan mm² (trích từ tên)'] + cot_moi
        elif pid == 'may_bien_ap_do_luong':
            cot_moi = ['Loại lõi (trích từ tên)'] + cot_moi
        cot_moi += ['Độ Tin Cậy Bổ Sung', 'Trạng Thái Kiểm Duyệt', 'Ghi Chú Kiểm Duyệt',
                   'Nghi Ngờ Phân Loại Sai', 'Trùng Mã/Tên Với Dòng Khác']
        header_moi = header + cot_moi
        idx_moi = {h: i for i, h in enumerate(header_moi)}

        # Sheet dữ liệu đầu tiên PHẢI tên đúng sheet_chuan để tool nhận diện
        # được (ràng buộc 1) — không chèn sheet nào khác trước nó.
        ws = wb_moi.create_sheet(sheet_chuan)
        ws.sheet_view.showGridLines = False
        for i, w in enumerate([max(10, min(40, len(str(h)) + 4)) for h in header_moi], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        # Header PHẢI ở dòng 1 (ràng buộc 2) — KHÔNG chèn dòng tiêu đề.
        _header_row(ws, 1, header_moi)

        dem_confirmed = dem_can_kiem = 0

        for r_i, row in enumerate(data, start=1):
            row = list(row)
            ten = s(row[idx['Tên hàng']]) or ''
            ma = s(row[idx['Mã hàng']])
            ghi_chu_list, gia_tri_them = [], {}

            if pid == 'cap_dien':
                if not s(row[idx['Tiết Diện (mm²)']]):
                    m = RE_TIET_DIEN.search(ten)
                    if m:
                        row[idx['Tiết Diện (mm²)']] = m.group(1).replace(' ', '')
                        ghi_chu_list.append('Tiết diện: trích từ tên (regex)')
                    else:
                        ghi_chu_list.append('Tiết diện: KHÔNG có trong cột gốc lẫn tên — để trống')
                if not s(row[idx['Cấp Điện Áp']]):
                    m = RE_CAP_DIEN_AP.search(ten)
                    if m:
                        row[idx['Cấp Điện Áp']] = f'{m.group(1)}{m.group(2)}'.replace(' ', '')
                        ghi_chu_list.append('Cấp điện áp: trích từ tên (regex)')
                    else:
                        ghi_chu_list.append('Cấp điện áp: KHÔNG có trong cột gốc lẫn tên — để trống')
                if not s(row[idx['Vật Liệu Ruột']]):
                    ghi_chu_list.append('[phụ] Vật liệu ruột: không có dữ liệu (outline đã coi field này là tuỳ chọn "nếu có")')

            elif pid == 'tu_va_thiet_bi_dong_cat':
                m = RE_CAP_DIEN_AP.search(ten)
                gia_tri_them['Cấp Điện Áp (trích từ tên)'] = f'{m.group(1)}{m.group(2)}'.replace(' ', '') if m else None
                m = RE_DONG_DINH_MUC.search(ten)
                gia_tri_them['Dòng Định Mức A (trích từ tên)'] = f'{m.group(1)}A' if m else None
                m = RE_TIET_DIEN.search(ten)
                gia_tri_them['Tiết Diện liên quan mm² (trích từ tên)'] = m.group(1).replace(' ', '') if m else None
                thieu = [k for k, v in gia_tri_them.items() if v is None]
                if thieu:
                    ghi_chu_list.append('Thiếu: ' + ', '.join(k.split(' (')[0] for k in thieu))
                ghi_chu_list.append('[phụ] Dòng cắt (kA): không có trong tên — outline đã coi field này là tuỳ chọn "khi có dữ liệu"')

            elif pid == 'may_bien_ap_do_luong':
                m = RE_LOAI_MBA.search(ten)
                gia_tri_them['Loại lõi (trích từ tên)'] = m.group(1) if m else None
                if not s(row[idx['Công Suất (kVA)']]):
                    m = RE_CONG_SUAT.search(ten)
                    if m:
                        row[idx['Công Suất (kVA)']] = m.group(1)
                        ghi_chu_list.append('Công suất: trích từ tên (regex)')
                    else:
                        ghi_chu_list.append('Công suất: KHÔNG có trong cột gốc lẫn tên — để trống')
                if not s(row[idx['Cấp Điện Áp (kV)']]):
                    m = RE_CAP_DIEN_AP.search(ten)
                    if m:
                        row[idx['Cấp Điện Áp (kV)']] = f'{m.group(1)}{m.group(2)}'.replace(' ', '')
                        ghi_chu_list.append('Cấp điện áp: trích từ tên (regex)')
                    else:
                        ghi_chu_list.append('Cấp điện áp: KHÔNG có trong cột gốc lẫn tên — để trống')
                ghi_chu_list.append('[phụ] Tổ đấu dây, tổn hao không tải/có tải: không có dữ liệu — cần catalogue hãng nếu muốn bổ sung')

            elif pid == 'phu_kien_va_khac':
                ghi_chu_list.append('Nhóm phụ kiện chưa có thông số kỹ thuật cấu trúc — cần catalogue hãng theo từng loại (đầu cose, sứ, vật tư...)')

            m = RE_TIEU_CHUAN.search(ten)
            gia_tri_them['Tiêu Chuẩn (trích từ tên)'] = m.group(0).upper().replace(' ', '') if m else None
            ghi_chu_phu = [] if m else ['Tiêu chuẩn: không có trong tên (field phụ, không bắt buộc)']

            co_thieu_cot_loi = any('KHÔNG có' in g or g.startswith('Thiếu') for g in ghi_chu_list)
            co_bo_sung_tu_ten = any('trích từ tên' in g for g in ghi_chu_list)
            if pid == 'phu_kien_va_khac':
                trang_thai, do_tin_cay = 'CAN_KIEM', 'Không có'
            elif co_thieu_cot_loi:
                trang_thai = 'CAN_KIEM'
                do_tin_cay = 'Trung bình' if co_bo_sung_tu_ten else 'Không có'
            else:
                trang_thai, do_tin_cay = 'CONFIRMED', 'Cao'
            ghi_chu_list = ghi_chu_list + ghi_chu_phu

            nghi_ngo = ''
            for tu_khoa, ly_do in TU_KHOA_NGHI_NGO.get(pid, []):
                if tu_khoa in ten.lower():
                    nghi_ngo = ly_do
                    break

            trung = []
            if ma and dem_ma.get(ma, 0) > 1:
                trung.append(f'Mã hàng "{ma}" xuất hiện {dem_ma[ma]} lần trong toàn hệ')
            if ten and dem_ten.get(ten, 0) > 1:
                trung.append(f'Tên hàng trùng {dem_ten[ten]} lần trong toàn hệ')
            trung_text = '; '.join(trung)

            hang_ghi = row + [gia_tri_them.get(c) for c in cot_moi[:-5]] + \
                [do_tin_cay, trang_thai, ' | '.join(ghi_chu_list), nghi_ngo, trung_text]

            r_ghi = r_i + 1  # header ở dòng 1 -> dữ liệu bắt đầu dòng 2 (ràng buộc 2+3)
            for c_i, v in enumerate(hang_ghi, 1):
                cell = ws.cell(row=r_ghi, column=c_i, value=v)
                cell.font = Font(name=F, size=9)
                cell.alignment = Alignment(vertical='top', wrap_text=(c_i > len(header)))
                cell.border = border
            mau = GREEN if trang_thai == 'CONFIRMED' else YELLOW
            ws.cell(row=r_ghi, column=idx_moi['Trạng Thái Kiểm Duyệt'] + 1).fill = PatternFill('solid', fgColor=mau)
            if nghi_ngo:
                ws.cell(row=r_ghi, column=idx_moi['Nghi Ngờ Phân Loại Sai'] + 1).fill = PatternFill('solid', fgColor=RED)
            if trung_text:
                ws.cell(row=r_ghi, column=idx_moi['Trùng Mã/Tên Với Dòng Khác'] + 1).fill = PatternFill('solid', fgColor=RED)

            if trang_thai == 'CONFIRMED':
                dem_confirmed += 1
            else:
                dem_can_kiem += 1

        ws.freeze_panes = 'A2'
        print(f'{pid}: {len(data)} dòng — CONFIRMED={dem_confirmed}, CAN_KIEM={dem_can_kiem}')

    _them_dashboard(wb_moi)
    wb_moi.save(duong_dan_ra)
    print(f'\n[OK] Đã lưu: {duong_dan_ra}')
    return duong_dan_ra


def _them_dashboard(wb):
    thong_tin_sheet = {}
    nhom_nguon_toan_he = set()
    for pid, (_, sheet_chuan, _) in TEN_FILE.items():
        ws = wb[sheet_chuan]
        header = [c.value for c in ws[1]]
        idx = {h: i for i, h in enumerate(header)}
        dong_cuoi = ws.max_row
        thong_tin_sheet[sheet_chuan] = {
            'dong_dau': 2, 'dong_cuoi': dong_cuoi,
            'cot_nhom': get_column_letter(idx['Nhóm Nguồn'] + 1),
            'cot_trang_thai': get_column_letter(idx['Trạng Thái Kiểm Duyệt'] + 1),
            'cot_nghi_ngo': get_column_letter(idx['Nghi Ngờ Phân Loại Sai'] + 1),
            'cot_trung': get_column_letter(idx['Trùng Mã/Tên Với Dòng Khác'] + 1),
        }
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[idx['Nhóm Nguồn']]:
                nhom_nguon_toan_he.add((sheet_chuan, r[idx['Nhóm Nguồn']]))

    # Dashboard PHẢI ở CUỐI danh sách sheet (ràng buộc 1) — không dùng index=0.
    ws = wb.create_sheet('00_Dashboard')
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([32, 30, 12, 12, 12, 14, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(row=r, column=1, value='LOTECH MASTER — DASHBOARD TỔNG QUAN (công thức sống, tự cập nhật)')
    c.font = Font(name=F, size=13, bold=True, color=WHITE)
    c.fill = PatternFill('solid', fgColor=NAVY)
    r += 2

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(row=r, column=1, value='TỔNG QUAN THEO NHÓM HÀNG')
    c.font = Font(name=F, size=11, bold=True, color=WHITE)
    c.fill = PatternFill('solid', fgColor='666666')
    r += 1
    for i, h in enumerate(['Profile', 'Tổng SP', 'CONFIRMED', 'CAN_KIEM', 'Nghi ngờ phân loại sai', 'Trùng mã/tên'], 1):
        cc = ws.cell(row=r, column=i, value=h)
        cc.font = Font(name=F, size=9, bold=True, color=NAVY)
        cc.fill = PatternFill('solid', fgColor=LBLUE)
        cc.border = border
    r += 1
    for pid, (_, sheet_chuan, nhan) in TEN_FILE.items():
        ti = thong_tin_sheet[sheet_chuan]
        rtt = f"'{sheet_chuan}'!{ti['cot_trang_thai']}{ti['dong_dau']}:{ti['cot_trang_thai']}{ti['dong_cuoi']}"
        rnn = f"'{sheet_chuan}'!{ti['cot_nghi_ngo']}{ti['dong_dau']}:{ti['cot_nghi_ngo']}{ti['dong_cuoi']}"
        rtr = f"'{sheet_chuan}'!{ti['cot_trung']}{ti['dong_dau']}:{ti['cot_trung']}{ti['dong_cuoi']}"
        ws.cell(row=r, column=1, value=nhan)
        ws.cell(row=r, column=2, value=f'=COUNTA({rtt})')
        ws.cell(row=r, column=3, value=f'=COUNTIF({rtt},"CONFIRMED")')
        ws.cell(row=r, column=4, value=f'=COUNTIF({rtt},"CAN_KIEM")')
        ws.cell(row=r, column=5, value=f'=COUNTIF({rnn},"?*")')
        ws.cell(row=r, column=6, value=f'=COUNTIF({rtr},"?*")')
        for i in range(1, 7):
            ws.cell(row=r, column=i).font = Font(name=F, size=9)
            ws.cell(row=r, column=i).border = border
        ws.cell(row=r, column=3).fill = PatternFill('solid', fgColor=GREEN)
        ws.cell(row=r, column=4).fill = PatternFill('solid', fgColor=YELLOW)
        r += 1
    r += 1

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(row=r, column=1, value='CHI TIẾT THEO NHÓM NGUỒN')
    c.font = Font(name=F, size=11, bold=True, color=WHITE)
    c.fill = PatternFill('solid', fgColor='666666')
    r += 1
    for i, h in enumerate(['Profile', 'Nhóm Nguồn', 'Tổng SP', 'CONFIRMED', 'CAN_KIEM'], 1):
        cc = ws.cell(row=r, column=i, value=h)
        cc.font = Font(name=F, size=9, bold=True, color=NAVY)
        cc.fill = PatternFill('solid', fgColor=LBLUE)
        cc.border = border
    r += 1
    nhan_map = {sc: nh for _, (_, sc, nh) in TEN_FILE.items()}
    for sheet_chuan, nhom in sorted(nhom_nguon_toan_he):
        ti = thong_tin_sheet[sheet_chuan]
        rnhom = f"'{sheet_chuan}'!{ti['cot_nhom']}{ti['dong_dau']}:{ti['cot_nhom']}{ti['dong_cuoi']}"
        rtt = f"'{sheet_chuan}'!{ti['cot_trang_thai']}{ti['dong_dau']}:{ti['cot_trang_thai']}{ti['dong_cuoi']}"
        ws.cell(row=r, column=1, value=nhan_map[sheet_chuan])
        ws.cell(row=r, column=2, value=nhom)
        ws.cell(row=r, column=3, value=f'=COUNTIF({rnhom},"{nhom}")')
        ws.cell(row=r, column=4, value=f'=COUNTIFS({rnhom},"{nhom}",{rtt},"CONFIRMED")')
        ws.cell(row=r, column=5, value=f'=COUNTIFS({rnhom},"{nhom}",{rtt},"CAN_KIEM")')
        for i in range(1, 6):
            ws.cell(row=r, column=i).font = Font(name=F, size=9)
            ws.cell(row=r, column=i).border = border
        r += 1


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description='Chuẩn hoá + làm giàu MASTER 4-profile (Ver 4.8)')
    ap.add_argument('--nguon', default='du_lieu_goc',
                    help='Thư mục chứa 4 file 0X_*.xlsx gốc, HOẶC 1 file .xlsx gộp sẵn 4 sheet '
                         '(VD MASTER_4_PROFILE_SAN_PHAM.xlsx) — tự nhận diện')
    ap.add_argument('--ra', default='MASTER_ENRICHED.xlsx', help='Đường dẫn file xuất ra')
    args = ap.parse_args()
    chuan_hoa(args.nguon, args.ra)
